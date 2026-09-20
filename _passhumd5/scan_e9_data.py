#!/usr/bin/env python3
"""HUM-D5 A3 E9 - data parity.
For every lesson with Data.xlsx / Data.csv or a config dataset: the numbers in
the workbook, the CSV, the lesson config dataset (read from the HTML SOURCE by
brace-matching window.CLASSIC_LESSON - the string-only text index has no
numbers) and the printed table in Pupil_Resources.html must agree; units and a
source line must be present in each. The blank/zero-input graph leg is a
Chromium check reported separately. Prints its search scope."""
import os, re, sys, json, csv, collections
from openpyxl import load_workbook
from lxml import html as LH
ROOT, OUT = sys.argv[1], sys.argv[2]
def config(t):
    i=t.find('window.CLASSIC_LESSON'); i=t.find('{',i); depth=0; instr=False; esc=False; q=''
    for j in range(i,len(t)):
        c=t[j]
        if instr:
            if esc: esc=False
            elif c=='\\': esc=True
            elif c==q: instr=False
        else:
            if c in '"\'': instr=True; q=c
            elif c=='{': depth+=1
            elif c=='}':
                depth-=1
                if depth==0: return json.loads(t[i:j+1])
from datetime import datetime
def num(s):
    s=str(s).strip().replace('−','-').replace(',','')
    try: return float(s)
    except: pass
    # an ISO timestamp in the config/CSV and an Excel datetime cell are the same instant
    for f in ('%Y-%m-%dT%H:%M:%S%z','%Y-%m-%d %H:%M:%S','%Y-%m-%dT%H:%M:%S'):
        try: return datetime.strptime(s,f).replace(tzinfo=None).timestamp()
        except ValueError: pass
    return None
def grid(rows): return [[str(c).strip() if c is not None else '' for c in r] for r in rows]
lessons=[]
for dp,_,fs in os.walk(ROOT):
    for fn in fs:
        if fn.endswith('_Lesson.html'): lessons.append(os.path.join(dp,fn))
out=[]; scope={'lessons':len(lessons),'with_dataset':0,'with_csv':0,'with_xlsx':0,'with_printed_table':0}
for lp in sorted(lessons):
    lid=os.path.basename(lp).replace('_Lesson.html',''); d=os.path.dirname(lp)
    t=open(lp,encoding='utf-8').read(); cfg=config(t)
    ds=cfg.get('dataset'); kind=cfg.get('kind')
    csvp=os.path.join(d,lid+'_Data.csv'); xlp=os.path.join(d,lid+'_Data.xlsx'); prp=os.path.join(d,lid+'_Pupil_Resources.html')
    has=dict(dataset=bool(ds),csv=os.path.exists(csvp),xlsx=os.path.exists(xlp))
    if not (has['dataset'] or has['csv'] or has['xlsx'] or kind=='graph'): continue
    for k in ('dataset','csv','xlsx'):
        if has[k]: scope['with_'+k]+=1
    rec={'lesson_id':lid,'kind':kind,'has':has,'findings':[]}
    cfg_tab=None
    if ds:
        cfg_tab=[[str(h) for h in ds.get('headers',[])]]+[[str(c) for c in r] for r in ds.get('rows',[])]
        rec['config']={'title':ds.get('title'),'units':ds.get('units'),'source':(ds.get('source') or '')[:120],'url':ds.get('url'),'limit':bool(ds.get('limit')),'rows':len(ds.get('rows',[]))}
        if not ds.get('units'): rec['findings'].append('config dataset has no units')
        if not ds.get('source'): rec['findings'].append('config dataset has no source line')
    csv_tab=None
    if has['csv']:
        rows=list(csv.reader(open(csvp,encoding='utf-8-sig')))
        meta={r[0]:r[1] for r in rows if len(r)>=2 and r[0] in ('Source','Limit','Units','Title')}
        body=[r for r in rows if r and r[0] not in ('Source','Limit','Units','Title')]
        csv_tab=grid(body); rec['csv']={'meta_keys':sorted(meta),'rows':len(body)-1}
        if 'Source' not in meta: rec['findings'].append('CSV has no Source line')
        if 'Units' not in meta and ds and ds.get('units') and not any(ds['units'] in ' '.join(r) for r in rows): rec['findings'].append('CSV carries no units text')
    xl_tab=None
    if has['xlsx']:
        # Workbook layout (measured): title / source / url / "Units: ..." / limit, then the
        # header row and data rows; every row is padded to the sheet width with empty cells.
        wb=load_workbook(xlp,data_only=True); ws=wb.worksheets[0]
        rows=[]
        for r in ws.iter_rows(values_only=True):
            r=[c for c in r]
            while r and r[-1] in (None,''): r.pop()
            if r: rows.append(r)
        hdr=None
        if ds and ds.get('headers'):
            for i,r in enumerate(rows):
                if [str(c).strip() for c in r]==[str(h) for h in ds['headers']]: hdr=i; break
        pre=[' '.join(str(c) for c in r) for r in rows[:hdr]] if hdr is not None else [' '.join(str(c) for c in r) for r in rows]
        units_line=next((x for x in pre if x.startswith('Units:')),None)
        source_line=next((x for x in pre if ds and ds.get('source') and x.startswith(ds['source'][:30])),None)
        url_line=next((x for x in pre if x.startswith('http')),None)
        if hdr is not None:
            end=len(rows)
            for i in range(hdr+1,len(rows)):          # the data block ends at the "Your enquiry" prompt block (single-cell rows)
                if len(rows[i])==1: end=i; break
            xl_tab=grid(rows[hdr:end]); rec['xlsx']={'sheets':wb.sheetnames,'header_row':hdr+1,'rows':len(xl_tab)-1,'units_line':units_line,'source_line':bool(source_line),'url_line':url_line}
            if not units_line: rec['findings'].append('XLSX has no Units line')
            elif ds.get('units') and units_line.replace('Units:','').strip()!=ds['units']: rec['findings'].append(f'XLSX units {units_line!r} != config {ds["units"]!r}')
            if not source_line: rec['findings'].append('XLSX has no source line matching the config dataset source')
            if ds.get('url') and url_line!=ds['url']: rec['findings'].append(f'XLSX url {url_line!r} != config {ds["url"]!r}')
        else:
            rec['xlsx']={'sheets':wb.sheetnames,'template':True,'first_rows':[' | '.join(str(c) for c in r)[:80] for r in rows[:3]]}
            rec['findings'].append('XLSX is an input template (no dataset header row) - nothing to compare' if not ds else 'XLSX has no row matching the config headers')
    pr_tab=None
    if os.path.exists(prp):
        doc=LH.fromstring(open(prp,encoding='utf-8').read())
        tabs=doc.xpath('//table')
        if tabs:
            scope['with_printed_table']+=1
            tb=tabs[0]; pr_tab=[[c.text_content().strip() for c in tr.xpath('./th|./td')] for tr in tb.xpath('.//tr')]
            ptxt=doc.text_content()
            rec['printed']={'tables':len(tabs),'rows':len(pr_tab)-1,'units_on_page':bool(ds and ds.get('units') and ds['units'] in ptxt),'source_on_page':bool(ds and ds.get('source') and ds['source'][:40] in ptxt)}
            if ds and ds.get('units') and ds['units'] not in ptxt: rec['findings'].append('printed page lacks the dataset units')
            if ds and ds.get('source') and ds['source'][:40] not in ptxt: rec['findings'].append('printed page lacks the source line')
    # numeric comparison: cell-by-cell on the header+rows grids
    def cmp(a,b,na,nb):
        if a is None or b is None: return
        if len(a)!=len(b) or any(len(x)!=len(y) for x,y in zip(a,b)):
            rec['findings'].append(f'{na} vs {nb}: shape {len(a)}x{len(a[0]) if a else 0} != {len(b)}x{len(b[0]) if b else 0}'); return
        for i,(x,y) in enumerate(zip(a,b)):
            for j,(p,q) in enumerate(zip(x,y)):
                np_,nq=num(p),num(q)
                if (np_ is not None and nq is not None and abs(np_-nq)>1e-9) or ((np_ is None or nq is None) and p.strip()!=q.strip()):
                    rec['findings'].append(f'{na} vs {nb} r{i}c{j}: {p!r} != {q!r}')
    cmp(cfg_tab,csv_tab,'config','csv'); cmp(cfg_tab,xl_tab,'config','xlsx'); cmp(csv_tab,xl_tab,'csv','xlsx'); cmp(cfg_tab,pr_tab,'config','printed')
    out.append(rec)
json.dump({'scope':scope,'lessons':out},open(OUT,'w'),indent=1,ensure_ascii=False)
print('SCOPE:',scope)
for r in out:
    print(f"{r['lesson_id']:16s} kind={r['kind']!s:9s} dataset={r['has']['dataset']} csv={r['has']['csv']} xlsx={r['has']['xlsx']} printed={'printed' in r} findings={len(r['findings'])}")
    for f in r['findings'][:8]: print('     -',f)
