#!/usr/bin/env python3
"""Editable teacher notes and exact source graphics for the P1 companions."""
import sys,json,datetime,zipfile,re
from pathlib import Path
sys.dont_write_bytecode=True
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from content import DATA,ROUTES,STAGES,TIMERS,ACCESS,SPACE,CAPTURE,CODES,CAVEAT
from build import ROOT,local_map,geo_map
def normalise_zip(p):
 with zipfile.ZipFile(p) as z:items={n:z.read(n) for n in z.namelist()}
 with zipfile.ZipFile(p,'w',zipfile.ZIP_DEFLATED,compresslevel=9) as z:
  for n,b in sorted(items.items()):
   if n=='docProps/core.xml':b=re.sub(rb'<dcterms:(created|modified)[^>]*>[^<]*</dcterms:\1>',lambda m:re.sub(rb'>[^<]*<',b'>2026-09-21T00:00:00Z<',m.group()),b)
   i=zipfile.ZipInfo(n,(2026,9,21,0,0,0));i.compress_type=zipfile.ZIP_DEFLATED;z.writestr(i,b)
for p,d in DATA.items():
 folder=ROOT/'packs'/f'HUM_Summer_1_{p}_Final'/p/'Summer_1/W01'
 doc=Document();s=doc.sections[0];s.page_width=Mm(210);s.page_height=Mm(297);s.top_margin=s.bottom_margin=Mm(18);s.left_margin=s.right_margin=Mm(20)
 for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
  st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
 doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.space_after=Pt(6)
 doc.styles['Title'].font.size=Pt(23)
 for st in doc.styles:
  for node in st.element.xpath('.//w:pBdr'):node.getparent().remove(node)
 for name,size in [('Heading 1',16),('Heading 2',13),('Heading 3',11)]:doc.styles[name].font.size=Pt(size)
 doc.core_properties.title=d['title']+' Teacher notes';doc.core_properties.author='Progress Schools';doc.core_properties.created=doc.core_properties.modified=datetime.datetime(2026,9,21)
 doc.add_paragraph(re.sub(r'[^\w\s]','',d['title'])+' Teacher notes','Title')
 doc.add_paragraph(f'{p} Humanities | Sum1 W1 | SoW slot 27 | 40 minutes in one session')
 doc.add_paragraph('Teach the exact weekly objective below using the nine-stage lesson. Select one route for each pupil, receive their response in an accessible form and agree a next step. These notes include the worked explanations, answers and evidence expectations.')
 doc.add_heading('Objective and evidence',1);doc.add_paragraph(d['objective']);doc.add_paragraph(d['evidence'])
 for x in d['success']:doc.add_paragraph(x,style='List Bullet')
 doc.add_heading('Preparation and access',1);doc.add_paragraph(d['prep']);doc.add_paragraph(ACCESS);doc.add_paragraph(d['misconception'])
 doc.add_heading('Vocabulary to teach first',1)
 for w,v,t in d['words']:doc.add_paragraph(f'{w} ({t}): {v}.')
 doc.add_paragraph(d['career'])
 doc.add_page_break();doc.add_heading('Teaching sequence',1)
 table=doc.add_table(rows=1,cols=3);table.style='Table Grid';table.rows[0].cells[0].text='Stage';table.rows[0].cells[1].text='Minutes';table.rows[0].cells[2].text='Teacher action and pupil task'
 for i,stage in enumerate(STAGES):
  cells=table.add_row().cells;cells[0].text=stage;cells[1].text=str(TIMERS[i]);cells[2].text=d['tasks'][i]
 for row in table.rows:
  for cell in row.cells:
   tcPr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   tcPr.append(borders)
 for cell in table.rows[0].cells:
  sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E5EDF3');cell._tc.get_or_add_tcPr().append(sh)
 for i,m in enumerate(d['models']):doc.add_heading('I Do worked explanation '+str(i+1),2);doc.add_paragraph(m)
 doc.add_heading('We Do checks',2)
 for c in d['checks']:doc.add_paragraph(c['question']);doc.add_paragraph(c['feedback'])
 doc.add_page_break();doc.add_heading('Arrival questions and answers',1)
 doc.add_paragraph('Use one route only. Hints sit beside questions in the lesson and pupil resources. Accept equivalent meaning in the pupil’s response mode.')
 for r in ROUTES:
  doc.add_heading(r.title(),2)
  for i,(h,q,hint,a) in enumerate(d['arrival'][r]):
   doc.add_paragraph(f'{i+1}. {q}');para=doc.add_paragraph();para.add_run('Answer: ').bold=True;para.add_run(a)
 doc.add_page_break();doc.add_heading('Independent task and exit',1)
 for i,r in enumerate(ROUTES):
  doc.add_heading(r.title(),2);doc.add_paragraph(d['independent'][i]);doc.add_paragraph('Exit: '+d['exit'][i][0]);doc.add_paragraph('Expected evidence: '+d['exit'][i][1])
 doc.add_heading('Use the feedback response loop',1)
 doc.add_paragraph('Provide Space, invite Voice, genuinely receive the response as Audience, then agree Influence. A button records the agreed step; it does not itself prove learning. Keep each stage’s response separate. In I Do stages the teacher models without a pupil-response panel. Use the next stage to check understanding.')
 doc.add_paragraph('Record the support used. If a pupil declines, consider regulation, communication, demand, relationship and timing; offer another route or revisit. Do not make completion of a digital panel a condition for access to the next task.')
 doc.add_heading('Accreditation caveat from the SoW',2);doc.add_paragraph(CAVEAT)
 doc.add_page_break();doc.add_heading('Policy wording and source checks',1)
 doc.add_paragraph('The quotations below use the supplied Feedback and Marking Policy 2025/2026 Pilot, Issue 1. Body section numbers are used because the contents page differs.')
 doc.add_heading('Space',2)
 for i,q in enumerate(SPACE):doc.add_paragraph('“'+q+'”');
 doc.add_paragraph('First five quotations: body §5, PDF page 8. Final quotation: body §14, PDF page 14.')
 doc.add_heading('Staff codes',2)
 doc.add_paragraph(' · '.join(c+' = '+v for c,v in CODES))
 doc.add_heading('EFL capture at Exit',2);doc.add_paragraph('“'+CAPTURE+'”');doc.add_paragraph('Feedback Policy body §14, PDF page 14. Use an individual response, not an undifferentiated group outcome.')
 doc.add_heading('Content sources',2)
 for title,url,note in d['sources']:doc.add_paragraph(title+'. '+note);doc.add_paragraph(url)
 footer=s.footer.paragraphs[0];footer.text=f'{p} Humanities | P1 preview | 21 September 2026 | '
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
 path=folder/'Teacher_Notes.docx';doc.save(path);normalise_zip(path)
 for kind,markup in [('local',local_map()),('uk',geo_map('uk')),('world',geo_map('world'))]:
  svg=re.search(r'<svg.*?</svg>',markup,re.S).group();svg=svg.replace('<svg ','<svg xmlns="http://www.w3.org/2000/svg" ',1);(ROOT/'qa'/f'{kind}.svg').write_text(svg)
print('Created 3 editable teacher notes and source map SVGs')
