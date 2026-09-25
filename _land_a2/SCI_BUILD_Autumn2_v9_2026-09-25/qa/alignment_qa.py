"""Read-only P16–P19 checks. No browser/keyboard/print execution."""
from pathlib import Path
from lxml import html
from docx import Document
import json,re,sys
import fitz
TARGET="Capture E against the pupil's own IEDP target on {EVIDENCE_APP}; the IEDP carries their EHCP-linked and Boxall targets, so one record serves every plan. Follow the pupil's IBSP where it differs from these routines."
def norm(t):return re.sub(r'\s+',' ',t).strip()
def cfg(d):return json.loads(next(s for s in d.xpath('//script/text()') if s.startswith('window.CLASSIC_LESSON')).split('=',1)[1].strip().rstrip(';'))
TERM_CODES={'Autumn 2':'A2','Spring 1':'SP1','Spring 2':'SP2','Summer 1':'SU1','Summer 2':'SU2'}
def expected_code(c):return f"{'SCI' if c.get('subject')=='Science' else 'HUM'}-{c['pathway']}-{TERM_CODES[c['term']]}-W{int(re.sub(r'\D','',str(c['week']))):02d}"
def audit(pack):
 science=bool(list(pack.rglob('LESSON_SOURCE.json')));items=list(pack.rglob('LESSON_SOURCE.json')) if science else list(pack.rglob('*_Lesson.html'));rows=[]
 for item in sorted(items):
  if science:
   c=json.loads(item.read_text());p=next(p for p in item.parent.glob('*.html') if p.name!='KNOWLEDGE_ORGANISER.html');staff=(item.parent/'TEACHER_NOTE.md').read_text()
  else:
   p=item;c=cfg(html.fromstring(p.read_text()));staff='\n'.join(x.text for x in Document(next(p.parent.glob('*Teacher_Notes.docx'))).paragraphs)
  e=c['planning_entry'];checks={}
  code=expected_code(c)
  checks['P16_P17_lesson_code_identical']=c.get('lesson_code')==code and c.get('book_header',{}).get('lesson_code')==code and e.get('lesson_code')==code
  checks['P17_evidence_code_instruction']=e.get('evidence_label_line','').startswith('Label evidence on {EVIDENCE_APP} with this code: '+code+'.') and all(x in e.get('evidence_label_line','') for x in ['spoken','practical','photographic']) and e['evidence_label_line'] in staff
  if science:checks['P17_source_code']=c.get('lesson_code')==code
  for hp in p.parent.glob('*.html'):
   if hp!=p and not any(x in hp.name for x in ['Knowledge_Organiser','Pupil_Resources','KNOWLEDGE_ORGANISER']):continue
   d=html.fromstring(hp.read_text());hdr=d.xpath('//*[@data-book-header="P16"]');s=norm(hdr[0].text_content()) if hdr else ''
   checks['P16_header_'+hp.name]=bool(hdr) and all(t in s for t in [c.get('subject','Humanities'),c['pathway'],c['term'],c['title'],code,'Date','Today I can','☐','VF','WS','NS','E','R']) and bool(d.xpath('//style[@id="p16-print-layer"]'))
  if not science:
   for pdf in list(p.parent.glob('*Pupil_Resources.pdf'))+list(p.parent.glob('*Knowledge_Organiser.pdf')):
    checks['P16_every_PDF_page_'+pdf.name]=all(all(t in pg.get_text() for t in [c['pathway'],c['term'],code,'Date','Today I can','Feedback:']) for pg in fitz.open(pdf))
   doc=Document(next(p.parent.glob('*Editable_Pack.docx')));checks['P16_editable_repeating_header']=all(all(t in ' '.join(p.text for p in s.header.paragraphs) for t in ['Date','Today I can','Feedback:',c['pathway'],code]) for s in doc.sections)
  checks['P17_exact_outcome_reference']=bool(e['sow_reference']) and e['exact_sow_outcome'] in staff and e['sow_reference'] in staff
  checks['P17_first_line_lesson_code']=bool(re.search(r'SLT planning and evaluation — P17\s+Lesson code: '+re.escape(code),staff))
  checks['P17_eight_stage_lines']=len(e['stages'])==8 and all(x in staff for x in e['stages'])
  checks['P17_exact_evidence_cell']=e.get('sow_evidence_reference','') in staff and bool(e.get('sow_evidence_reference'))
  checks['P17_evidence_and_if_behind']=all(e[k] in staff for k in ['sow_evidence_line','evidence_in_books']) and all(e['if_behind'].values())
  checks['P17_blank_evaluation']=all(e['staff_evaluation'].values()) and all(k in e['staff_evaluation'] for k in ['date_taught','delivered_as_planned','criteria_met','adaptations_made_and_why','change_from_sow_reason_and_coverage_impact','evidence_banked_on_{EVIDENCE_APP}'])
  brief=(p.parent/'TA_BRIEF.md').read_text();checks['P18_five_target_areas']=len(e['target_links'])==5 and all(v in staff and v in brief for v in e['target_links'].values());checks['P18_recording_line']=TARGET in staff and TARGET in brief
  d=html.fromstring(p.read_text())
  for node in d.xpath('//script|//style|//*[@data-mbm-guide]|//*[@id="ta-dialog"]'):
   if node.getparent() is not None:node.getparent().remove(node)
  checks['P18_no_target_systems_on_pupil_page']=not any(x in d.text_content() for x in ['IEDP','EHCP-linked','Boxall','IBSP','{EVIDENCE_APP}'])
  rows.append({'id':c['id'],'checks':checks,'failed':[k for k,v in checks.items() if not v]})
 unit=all((pack/f'UNIT_RATIONALE.{s}').exists() for s in ['html','docx'])
 return {'status':'PASS' if unit and all(not r['failed'] for r in rows) else 'FAIL','P19_unit_files':unit,'checks_passed':sum(sum(r['checks'].values()) for r in rows)+int(unit),'checks_total':sum(len(r['checks']) for r in rows)+1,'lessons':rows,'browser':'NOT RUN','keyboard':'NOT RUN','print':'NOT RUN','limits':'PDF text/page checks and OOXML headers are static artifact checks; browser pagination still belongs to Claude.'}
if __name__=='__main__':
 r=audit(Path(sys.argv[1]));print(json.dumps(r,ensure_ascii=False,indent=2));raise SystemExit(r['status']!='PASS')
