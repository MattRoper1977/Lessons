from pathlib import Path
import json, sys, subprocess, shutil, hashlib, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from visuals import render_svg

BASE=Path(__file__).resolve().parents[1]
NODE='/opt/codex/runtimes/codex-primary-runtime/dependencies/node/bin/node'
PY='/opt/codex/runtimes/codex-primary-runtime/dependencies/python/bin/python3'
COLORS={'BUILD':'4E7A9B','GROW':'3F7D6E','LAUNCH':'7A5C9E'}

def clean(v):
 return str(v).replace('\u2011','-').replace('\u2013','-').replace('\u2014',' - ')

def ptext(doc,txt,style=None,bold=False,size=None):
 p=doc.add_paragraph(style=style); r=p.add_run(clean(txt));r.bold=bold
 if size:r.font.size=Pt(size)
 return p

def base(d,pupil=False):
 doc=Document();s=doc.sections[0];s.page_width=Inches(8.2677);s.page_height=Inches(11.6929)
 s.top_margin=Inches(.57);s.bottom_margin=Inches(.72);s.left_margin=s.right_margin=Inches(.68)
 s.header_distance=s.footer_distance=Inches(.25)
 st=doc.styles['Normal'];st.font.name='Arial';st.font.size=Pt(13 if pupil else 10.5)
 st.paragraph_format.space_after=Pt(6 if pupil else 4);st.paragraph_format.line_spacing=1.12
 for n,sz in [('Title',24),('Subtitle',12),('Heading 1',18),('Heading 2',14),('Heading 3',12)]:
  st=doc.styles[n];st.font.name='Arial';st.font.size=Pt(sz);st.font.color.rgb=RGBColor(0,0,0)
  st.paragraph_format.space_before=Pt(8);st.paragraph_format.space_after=Pt(5)
  st.paragraph_format.keep_with_next=True
 footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 r=footer.add_run(f"MADE BY MATT  |  {d['pathway']} SCIENCE  |  ");r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('596578')
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 for st in doc.styles:
  for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 doc.core_properties.title=d['title'];doc.core_properties.author='Made by Matt';doc.core_properties.subject=f"{d['pathway']} Science teaching pack"
 return doc

def shade(cell,color):
 sh=OxmlElement('w:shd');sh.set(qn('w:fill'),color);cell._tc.get_or_add_tcPr().append(sh)

def table(doc,headers,rows,pupil=False):
 if not headers:return
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 short={'task','card','slide','minutes','stage','trial','test'}
 weights=[.7 if str(h).lower() in short else 2.1 for h in headers]
 widths=[6.90*x/sum(weights) for x in weights]
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for c,x in zip(t.rows[0].cells,headers):c.text=clean(x);shade(c,'DBE8F0')
 trPr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
 for row in rows:
  cells=t.add_row().cells
  for c,x in zip(cells,row):c.text=clean(x)
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();keep=OxmlElement('w:cantSplit');trpr.append(keep)
  if ri and pupil and any(not str(x).strip() for x in rows[ri-1]):
   row.height=Inches(.5);row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci])
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=c._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');bd.append(e)
   tcpr.append(bd);mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
   tcpr.append(mar)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.04
    if len(rows)<=8 and ri<len(t.rows)-1:p.paragraph_format.keep_with_next=True
    for r in p.runs:r.font.size=Pt(11.5 if pupil else 9.5);r.bold=(ri==0)
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=.2
 return t

def visual(doc,v,out,name):
 cache=BASE/'build'/'figures';cache.mkdir(exist_ok=True)
 uid=hashlib.sha256(json.dumps(v,sort_keys=True).encode()).hexdigest()[:12]
 svg=cache/(uid+'.svg');png=cache/(uid+'.png')
 if not png.exists():
  svg.write_text(render_svg(v,uid),encoding='utf8')
  subprocess.run([NODE,str(BASE/'build'/'raster.mjs'),str(svg),str(png)],check=True,capture_output=True)
 r=doc.add_paragraph().add_run();pic=r.add_picture(str(png),width=Inches(6.85))
 pic._inline.docPr.set('descr',v.get('description',v.get('title','Scientific diagram')))
 # Retain the editable diagram SVG alongside the pupil resources and PowerPoint.
 dest=out/'Diagrams';dest.mkdir(exist_ok=True);shutil.copyfile(svg,dest/(name+'.svg'))

def response(doc,prompt,lines=2,line_height=24):
 ptext(doc,prompt,bold=True)
 if re.search(r'\b(sketch|draw)\b',prompt,re.I):
  p=doc.add_paragraph();p.paragraph_format.space_after=Pt(line_height*lines);p.paragraph_format.line_spacing=Pt(1)
  p.add_run('\u00a0').font.size=Pt(1)
  return
 for _ in range(lines):
  p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
  p.paragraph_format.space_before=Pt(0);p.paragraph_format.keep_with_next=False
  r=p.add_run('_____________________________________________________________________');r.font.size=Pt(11);r.font.color.rgb=RGBColor.from_string('8792A1')
  p.paragraph_format.line_spacing=Pt(line_height)

def pupil_doc(d,out):
 doc=base(d,True)
 for idx,page in enumerate(d['pupilPages']):
  p=ptext(doc,f"{d['pathway']} SCIENCE  /  {d['sow']['term']}  /  {d['sow']['week']}",size=10)
  if idx:p.paragraph_format.page_break_before=True
  ptext(doc,page['title'],'Title')
  if not idx:
   ptext(doc,d['title'],bold=True,size=14)
   ptext(doc,'Name __________________________    Date ______________',size=11)
  if page.get('intro'):ptext(doc,page['intro'])
  for bi,b in enumerate(page.get('blocks',[])):
   typ=b['type']
   if typ=='text':ptext(doc,b['text'])
   elif typ=='visual':visual(doc,b['visual'],out,f'Pupil_{idx+1}_{bi+1}')
   elif typ=='table':table(doc,b['headers'],b['rows'],True)
   elif typ=='response':response(doc,b['prompt'],b.get('lines',2),page.get('responseLineHeight',24))
   elif typ=='choices':
    ptext(doc,b['prompt'],bold=True)
    for j,c in enumerate(b['choices']):ptext(doc,f"{chr(65+j)}. {c}")
   elif typ=='cards':
    cards=b['cards'];t=doc.add_table(rows=0,cols=2);t.autofit=False
    for col in t.columns:col.width=Inches(3.45)
    for j in range(0,len(cards),2):
     row=t.add_row();trPr=row._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
     for k,c in enumerate(row.cells):
      if j+k>=len(cards):continue
      x=cards[j+k];c.text='';p=c.paragraphs[0];p.add_run(clean(x['title'])).bold=True
      p=c.add_paragraph(clean(x['text']));p.paragraph_format.space_after=Pt(10)
      for pp in c.paragraphs:
       for rr in pp.runs:rr.font.size=Pt(12)
      shade(c,'F4F6F8')
    doc.add_paragraph()
   else:raise ValueError(typ)
 doc.save(out/'Pupil.docx')

def heading_paras(doc,title,items):
 if not items:return
 ptext(doc,title,'Heading 2')
 if isinstance(items,str):ptext(doc,items)
 else:
  for it in items:ptext(doc,it)

def hyperlink(doc,text,url):
 p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
 h=OxmlElement('w:hyperlink');rid=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True);h.set(qn('r:id'),rid)
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'1F6B4A');pr.append(co);r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)

def teacher_doc(d,out):
 doc=base(d)
 ptext(doc,d['pathway']+' SCIENCE  /  '+d['sow']['term']+'  /  '+d['sow']['week'],size=10)
 ptext(doc,d['title'],'Title');ptext(doc,'Teacher guide and worked answers','Subtitle')
 ptext(doc,d['objective'],bold=True,size=12)
 ptext(doc,'40 minutes. Use Lesson.html for interactive teaching or Lesson.pptx for editable slides. Slides.pdf is the static alternative. Select one pupil response route and print the shared evidence it requires. Detailed answers below are for staff.')
 heading_paras(doc,'Scheme of work',[f"{d['sow']['workbook']} | {d['sow']['sheet']} | {d['sow']['cells']}",d['sow']['exactOutcome'],d['sow']['alignment']])
 heading_paras(doc,'Prior learning',d['priorLearning']);heading_paras(doc,'Success evidence',d['success'])
 heading_paras(doc,'Equipment',d['equipment']);heading_paras(doc,'Preparation',d['preparation'])
 heading_paras(doc,'Practical care',d['safety']);heading_paras(doc,'Ready to teach alternative',d['practicalAlternative'])
 ptext(doc,'Teaching sequence','Heading 1')
 clock=0;timing=[]
 for i,s in enumerate(d['slides']):
  m=s.get('minutes',0);timer=f'{clock}-{clock+m}' if m else 'Within stage';clock+=m
  timing.append([str(i+1),s['stage'],timer,s['title']])
 table(doc,['Slide','Stage','Minutes','Focus'],timing)
 ptext(doc,'Times are a suggested 40-minute route. Slides marked Within stage share the preceding stage allocation. The optional HTML timer starts only when selected.')
 for i,s in enumerate(d['slides']):
  ptext(doc,f"{i+1}  {s['stage']}  {s['title']}",'Heading 2');ptext(doc,s['notes'])
  if s.get('reveal'):ptext(doc,'Reveal: '+s['reveal'])
  if s.get('options'):
   for j,o in enumerate(s['options']):ptext(doc,f"{chr(65+j)}. {o['text']} - {o['feedback']}")
 ptext(doc,'Answers and assessment','Heading 1')
 for a in d['answers']:
  ptext(doc,a['label'],'Heading 2');ptext(doc,a['answer'])
 heading_paras(doc,'Teaching and assessment notes',d['teacherNotes'])
 ptext(doc,'Access and responsive teaching','Heading 1')
 for k,v in d['access'].items():heading_paras(doc,k.title(),v)
 ptext(doc,'Misconceptions to check','Heading 1')
 for m in d['misconceptions']:
  p=ptext(doc,m['claim'],bold=True);p.paragraph_format.keep_with_next=True
  ptext(doc,m['correction']);ptext(doc,'Check: '+m['check'])
 table(doc,['Word','Meaning'],[[v['term'],v['meaning']] for v in d['vocabulary']])
 ptext(doc,'Scientific references','Heading 1')
 for src in d['sources']:
  hyperlink(doc,src['title'],src['url']);ptext(doc,src.get('note',''))
 doc.save(out/'Teacher.docx')

def render(d,out):
 records=[]
 for name in ['Pupil','Teacher']:
  q=BASE/'qa'/'docs'/d['id']/name
  if q.exists():shutil.rmtree(q)
  q.mkdir(parents=True,exist_ok=True)
  run=subprocess.run([PY,'/root/.codex/skills/builtins/documents/render_docx.py',str(out/(name+'.docx')),'--output_dir',str(q),'--emit_pdf','--dpi','105'],capture_output=True,text=True)
  if run.returncode:raise RuntimeError(run.stdout+run.stderr)
  shutil.copyfile(q/(name+'.pdf'),out/(name+'.pdf'))
  records.append({'lesson':d['id'],'file':name,'pages':len(list(q.glob('page-*.png'))),'path':str(q)})
 return records

if __name__=='__main__':
 files=[Path(x) for x in sys.argv[1:] if not x.startswith('--')] or sorted((BASE/'content').glob('*.json'))
 for f in files:
  source_bytes=f.read_bytes();source_sha=hashlib.sha256(source_bytes).hexdigest();d=json.loads(source_bytes);out=BASE/'output'/d['id'];out.mkdir(exist_ok=True,parents=True)
  diagrams=out/'Diagrams';diagrams.mkdir(exist_ok=True)
  for si,slide in enumerate(d['slides']):
   if slide.get('visual'):(diagrams/f'Diagram_slide_{si+1:02}.svg').write_text(render_svg(slide['visual'],d['id']+str(si)),encoding='utf8')
  pupil_doc(d,out);teacher_doc(d,out)
  if '--render' in sys.argv:
   record={'source_sha256':source_sha,'outputs':render(d,out)}
   (BASE/'qa'/'docs'/d['id']/'manifest.json').write_text(json.dumps(record,indent=2))
   print(json.dumps(record),flush=True)
  else:print('Created '+d['id'],flush=True)
