"""CX2 §3 · align the LAUNCH W4L1 Diffusion native files with the canonical lesson.

Reads the exact reviewed pack files from the git base recorded below and writes:

  * Week_4/W4L1_Diffusion.pptx   three arrival route slides, two knowledge-organiser
                                 slides, stage minutes as the lesson (4/1/4/4/3/5/15/4),
                                 one optional help prompt on We Do 2, a two-check exit
                                 per route; every new shape is a deep copy of a reviewed
                                 shape so the deck design is unchanged; no day or time
                                 wording remains (the lesson is one plan for every group)
  * Week_4/W4L1_Diffusion.pdf    rendered from that deck by LibreOffice
  * Week_4/W4L1_Worksheet.docx   the reviewed sheet plus arrival routes, organiser and
                                 exit routes; day and period wording removed
  * Week_4/W4L1_Worksheet.pdf    rendered from that document
  * Teacher_Guide_And_Answers.docx / .pdf   the W4L1 section gains the arrival and exit
                                 answers and the feedback card (Feedback & Marking Policy
                                 Issue 1 codes; no product name)
  * Pupil_Worksheets.pdf         the fifteen worksheet PDFs in order, W4L1 replaced

Content comes from the same JSON that authored the screen: W4L1_ARRIVAL.json,
W4L1_ORGANISER_EXIT.json and W4L1_FEEDBACK.json. Re-runnable: every output is
rebuilt from the git base, never from a previous output.

    python3 tools/launch_resources/author_w4l1_natives.py [--root REPO]
"""
from copy import deepcopy
from pathlib import Path
import argparse, io, json, re, shutil, subprocess, tempfile

from docx import Document
from docx.enum.text import WD_BREAK
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

BASE = '3a14e9c4d57832fed87ed588ef6f0ae4fd3d6064'   # Lessons main after #539; the pack bytes at review
HERE = Path(__file__).resolve().parent
PACK = 'Science_Teesside/Teaching_Packs/LAUNCH/'
DECK = PACK + 'Week_4/W4L1_Diffusion.pptx'
SHEET = PACK + 'Week_4/W4L1_Worksheet.docx'
GUIDE = PACK + 'Teacher_Guide_And_Answers.docx'
ARRIVAL = json.loads((HERE / 'W4L1_ARRIVAL.json').read_text())
KO = json.loads((HERE / 'W4L1_ORGANISER_EXIT.json').read_text())
FB = json.loads((HERE / 'W4L1_FEEDBACK.json').read_text())
EMU = 914400
A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'
ROUTES = ['supported', 'standard', 'stretch']
# The canonical lesson's stage minutes, by the deck's own stage names.
MINUTES = {'Arrival': 4, 'Starter': 1, 'I Do': 4, 'We Do': 4, 'I Do 2': 3, 'We Do 2 · Task': 5, 'Independent Task': 15, 'Exit': 4}
PLAN = 'Lesson 1 of Week 4, one 40-minute plan for every group'
DAYS = re.compile(r'\b(Monday|Tuesday|Wednesday|Thursday|Friday|P[12]\b|\d{2}:\d{2})')
REVISION = ('Revision 15 September 2026: aligned with the canonical lesson (four-question arrival routes on slides 1–3, '
            'knowledge organiser on slides 5–6, stage minutes 4/1/4/4/3/5/15/4 as the lesson, one optional help prompt '
            'on slide 10, two-check exit routes on slide 12). Pupil worksheet pages 3–6 carry the arrival routes, '
            'organiser and exit routes; staff answers and the feedback card are in Teacher_Guide_And_Answers.')


def baseline_bytes(root, rel):
    return subprocess.check_output(['git', '-C', str(root), 'show', BASE + ':' + rel])


# ----------------------------------------------------------------------------- pptx helpers
# The same reviewed-shape discipline as tools/grow_resources/author_w3_slides.py
# (Friction): copies of existing shapes, creation ids stripped, ids renumbered.
def strip_creation_ids(element):
    for ext in list(element.iter('{%s}extLst' % A_NS)):
        parent = ext.getparent()
        if parent is not None and parent.tag == qn('p:cNvPr'):
            parent.remove(ext)


def renumber(slide):
    next_id = 2
    for el in slide.shapes._spTree.iter(qn('p:cNvPr')):
        if el.getparent().getparent().tag == qn('p:nvGrpSpPr'):
            continue
        el.set('id', str(next_id)); next_id += 1


def clone_slide(prs, src):
    new = prs.slides.add_slide(src.slide_layout)
    for shape in list(new.shapes):
        shape._element.getparent().remove(shape._element)
    src_bg = src._element.cSld.find(qn('p:bg'))
    if src_bg is not None:
        new._element.cSld.insert(0, deepcopy(src_bg))
    for el in src.shapes._spTree:
        if el.tag in (qn('p:sp'), qn('p:cxnSp'), qn('p:grpSp')):
            copy = deepcopy(el); strip_creation_ids(copy); new.shapes._spTree.append(copy)
        else:
            assert el.tag not in (qn('p:pic'), qn('p:graphicFrame')), 'clone source carries a picture or table'
    # hyperlinks on copied shapes point at relationship ids of the source part
    rids = {}
    for hl in new.shapes._spTree.iter(qn('a:hlinkClick')):
        rid = hl.get(qn('r:id'))
        if rid and rid not in rids:
            target = src.part.rels[rid]
            rids[rid] = (new.part.rels.get_or_add_ext_rel(target.reltype, target.target_ref) if target.is_external
                         else new.part.relate_to(target.target_part, target.reltype))
        if rid:
            hl.set(qn('r:id'), rids[rid])
    renumber(new)
    return new


def move_slide(prs, old_index, new_index):
    lst = prs.slides._sldIdLst; items = list(lst); el = items[old_index]
    lst.remove(el); lst.insert(new_index, el)


def set_text(shape, text):
    body = shape.text_frame._txBody
    for extra in body.findall(qn('a:p'))[1:]:
        body.remove(extra)
    paragraph = shape.text_frame.paragraphs[0]
    runs = paragraph.runs
    runs[0].text = text
    for extra in runs[1:]:
        extra._r.getparent().remove(extra._r)


def set_paragraphs(shape, lines):
    body = shape.text_frame._txBody
    paras = body.findall(qn('a:p'))
    template = deepcopy(paras[0])
    for p in paras:
        body.remove(p)
    for line in lines:
        p = deepcopy(template)
        runs = p.findall(qn('a:r'))
        runs[0].find(qn('a:t')).text = line
        for extra in runs[1:]:
            p.remove(extra)
        body.append(p)


def set_size(shape, sz, bold=None):
    for rpr in shape._element.iter(qn('a:rPr'), qn('a:defRPr')):
        rpr.set('sz', str(sz))
        if bold is not None:
            rpr.set('b', '1' if bold else '0')


def place(shape, x, y, cx, cy):
    shape.left, shape.top, shape.width, shape.height = int(x * EMU), int(y * EMU), int(cx * EMU), int(cy * EMU)


def add_from(slide, template_shape, name, x, y, cx, cy, text=None, lines=None, sz=None, bold=None, anchor=None):
    el = deepcopy(template_shape._element); strip_creation_ids(el)
    for hl in list(el.iter(qn('a:hlinkClick'))):
        hl.getparent().remove(hl)
    if anchor is not None:
        el.find(qn('p:txBody')).find(qn('a:bodyPr')).set('anchor', anchor)
    slide.shapes._spTree.append(el)
    shape = slide.shapes[-1]
    shape.name = name
    place(shape, x, y, cx, cy)
    if lines is not None:
        set_paragraphs(shape, lines)
    elif text is not None:
        set_text(shape, text)
    if sz is not None:
        set_size(shape, sz, bold)
    renumber(slide)
    return shape


def remove_shapes(slide, names):
    for shape in list(slide.shapes):
        if shape.name in names:
            shape._element.getparent().remove(shape._element)


def nav_inactive(slide, name):
    """Give a cloned slide's highlighted footer button the plain state of its neighbour."""
    plain = next(s for s in slide.shapes if s.name == 'goto:6')
    active = next(s for s in slide.shapes if s.name == name)
    for tag in ('a:solidFill', 'a:ln'):
        old = active._element.find(qn('p:spPr')).find(qn(tag)); new = plain._element.find(qn('p:spPr')).find(qn(tag))
        if old is not None and new is not None:
            old.getparent().replace(old, deepcopy(new))
    for rpr_old, rpr_new in zip(active._element.iter(qn('a:rPr')), plain._element.iter(qn('a:rPr'))):
        rpr_old.getparent().replace(rpr_old, deepcopy(rpr_new))


def at(slide, left, top):
    for shape in slide.shapes:
        if shape.left == left and shape.top == top and shape.has_text_frame:
            return shape
    raise KeyError((left, top))


def header(slide):
    return at(slide, 8382000, 428625)


def title(slide):
    return at(slide, 609600, 866775)


def notes_text(slide):
    return slide.notes_slide.notes_text_frame.text


def set_notes(slide, paragraphs):
    slide.notes_slide.notes_text_frame.text = '\n'.join(paragraphs)


def replace_notes(slide, replacements):
    text = notes_text(slide)
    for old, new in replacements.items():
        assert old in text, (old, text[:160])
        text = text.replace(old, new)
    set_notes(slide, text.split('\n'))


def check_no_day_words(prs):
    for index, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.has_text_frame:
                assert not DAYS.search(shape.text_frame.text), (index, shape.name, shape.text_frame.text)
        assert not DAYS.search(notes_text(slide)), (index, notes_text(slide)[:200])


# ----------------------------------------------------------------------------- deck
def build_deck(root):
    prs = Presentation(io.BytesIO(baseline_bytes(root, DECK)))
    slides = list(prs.slides)
    assert len(slides) == 17
    arrival = slides[0]
    cue, text, strip, line = at(arrival, 666750, 3495675), at(arrival, 1162050, 3495675), at(arrival, 666750, 3076575), at(arrival, 666750, 1819275)
    templates = {'cue': deepcopy(cue._element), 'text': deepcopy(text._element), 'strip': deepcopy(strip._element), 'line': deepcopy(line._element)}
    class T:  # template shapes as objects with ._element
        pass
    tpl = {}
    for key, el in templates.items():
        obj = T(); obj._element = el; tpl[key] = obj

    # ---- stage minutes as the lesson, on the stage header and in the notes
    for slide in slides:
        head = header(slide).text_frame.text
        m = re.match(r'(.+?)\s+·\s+(\d+) min$', head)
        if not m:
            continue
        stage, old = m.group(1), int(m.group(2))
        new = MINUTES[stage]
        set_text(header(slide), f'{stage}  ·  {new} min')
        replace_notes(slide, {f'Stage: {stage}, {old} minutes. Monday / Wednesday / Friday P1, 09:30–10:10.':
                              f'Stage: {stage}, {new} minutes. {PLAN}.'})

    # ---- arrival: the reviewed slide becomes the Supported route; two clones follow it
    for shape in list(arrival.shapes):
        if shape.shape_type == 13:   # the retrieval picture makes way for the four questions
            shape._element.getparent().remove(shape._element)
    remove_shapes(arrival, [s.name for s in arrival.shapes if s.name in ('visual-ground', 'goto:8')
                            or (s.has_text_frame and s.top in (1819275, 3076575, 3495675, 4391025))])
    route_slides = []
    tail = notes_text(arrival).split('\n')[2:]   # core explanation onwards; the retired retrieval answers go
    for route in ROUTES:
        data = ARRIVAL['routes'][route]
        slide = arrival if route == 'supported' else clone_slide(prs, arrival)
        route_slides.append(slide)
        set_text(title(slide), f'Arrival · {data["label"]} route')
        add_from(slide, tpl['strip'], 'strip', 0.729, 1.90, 12.0, 0.36,
                 text='FOUR QUICK QUESTIONS   /   POINT, SAY, SIGN, DRAW OR WRITE   /   YOU MAY PASS', sz=1200, bold=True)
        y = 2.40
        for i, question in enumerate(data['questions']):
            add_from(slide, tpl['cue'], f'cue-{i}', 0.729, y, 0.47, 0.42, text=str(i + 1))
            add_from(slide, tpl['text'], f'question-{i}', 1.27, y, 11.4, 0.78, text=question, sz=1800)
            y += 0.84
        add_from(slide, tpl['line'], 'access', 0.729, 5.80, 11.9, 0.95, text=ARRIVAL['access'] + ' ' + ARRIVAL['response'], sz=1200)
        set_notes(slide, [
            f'Stage: Arrival, 4 minutes shared across the three route slides; show the route you have chosen. {PLAN}. '
            'Use reveals within this slot; appendix slides are not added time.',
            ARRIVAL['staff'] + ' Pupil worksheet page 3 carries all three routes; staff answers are in the teacher guide.',
            f'[Staff answers — {data["label"]}; do not display]',
            *[f'{i + 1}. {answer}' for i, answer in enumerate(data['answers'])],
            *tail])
    count = len(prs.slides)
    move_slide(prs, count - 2, 1); move_slide(prs, count - 1, 2)

    # ---- knowledge organiser, two reference slides after the Starter
    slides = list(prs.slides)
    check = slides[10]   # 'Retrieval and prediction check' (reviewed slide 9)
    assert title(check).text_frame.text == 'Retrieval and prediction check'
    def ko_slide(name):
        slide = clone_slide(prs, check)
        for s in list(slide.shapes):
            if s.name == 'prediction-answer' or (s.name == 'goto:1' and s.top == 5581650) or (s.has_text_frame and s.top in (1885950, 2647950, 4619625)):
                s._element.getparent().remove(s._element)
        nav_inactive(slide, 'goto:12')
        set_text(header(slide), 'Knowledge organiser')
        set_text(title(slide), name)
        set_notes(slide, [f'Knowledge organiser — 0 minutes (reference). {PLAN}. Return to this slide at any point; it matches '
                          'pupil worksheet pages 4–5 and the on-screen organiser. It is not an extra stage.',
                          *notes_text(check).split('\n')[1:]])
        return slide
    ko1 = ko_slide('Diffusion knowledge organiser (1 of 2)')
    add_from(ko1, tpl['line'], 'objective', 0.729, 1.85, 11.9, 0.45, text='Objective: ' + KO['objective'], sz=1600, bold=True)
    add_from(ko1, tpl['strip'], 'words-header', 0.729, 2.38, 7.4, 0.34, text='KEY WORDS', sz=1200, bold=True)
    add_from(ko1, tpl['text'], 'words', 0.729, 2.75, 7.4, 3.9, anchor='t',
             lines=[f'{word} ({tier}) — {meaning}' for word, tier, meaning in KO['vocabulary']], sz=1350)
    add_from(ko1, tpl['strip'], 'success-header', 8.35, 2.38, 4.3, 0.34, text='SUCCESS CRITERIA', sz=1200, bold=True)
    add_from(ko1, tpl['text'], 'success', 8.35, 2.75, 4.3, 3.9, anchor='t',
             lines=[f'{i + 1}. {item}' for i, item in enumerate(KO['success'])], sz=1300)
    ko2 = ko_slide('Diffusion knowledge organiser (2 of 2)')
    add_from(ko2, tpl['strip'], 'ideas-header', 0.729, 1.85, 11.9, 0.34, text='KEY IDEAS', sz=1200, bold=True)
    add_from(ko2, tpl['text'], 'ideas', 0.729, 2.22, 11.9, 2.35, anchor='t', lines=KO['facts'], sz=1450)
    add_from(ko2, tpl['strip'], 'model-header', 0.729, 4.60, 11.9, 0.34, text='THE MODEL', sz=1200, bold=True)
    add_from(ko2, tpl['text'], 'model', 0.729, 4.95, 11.9, 0.62, text=KO['model_caption'], sz=1200)
    add_from(ko2, tpl['strip'], 'examples-header', 0.729, 5.60, 11.9, 0.34, text='DIFFUSION OR NOT?', sz=1200, bold=True)
    add_from(ko2, tpl['text'], 'examples', 0.729, 5.95, 11.9, 0.95, anchor='t', lines=KO['examples'], sz=1150)
    count = len(prs.slides)
    move_slide(prs, count - 2, 4); move_slide(prs, count - 1, 5)

    # ---- We Do 2 task slide: one optional help prompt line, no response field
    slides = list(prs.slides)
    task = slides[9]
    assert title(task).text_frame.text == 'Count both directions'
    add_from(task, tpl['line'], 'help-prompt', 0.729, 6.27, 7.75, 0.60,
             text=FB['prompt_title'] + ' ' + FB['prompt_body'], sz=1000)
    replace_notes(task, {'Run the printed activity or this board; one route is enough.':
                         'Run the printed activity or this board; one route is enough. The help line is the lesson’s one optional '
                         'prompt: a pupil may point, say, sign or show an idea, question or change; the adult listens, explains what '
                         'can change and checks whether it helped. Passing is allowed; no extra record.'})

    # ---- exit: two small checks per route, answers in notes and the teacher guide
    exit_slide = slides[11]
    assert title(exit_slide).text_frame.text == 'What changed in your thinking?'
    remove_shapes(exit_slide, [s.name for s in exit_slide.shapes if s.has_text_frame and s.top in (2028825, 3629025, 4286250, 4943475)])
    set_text(title(exit_slide), 'Before you go · two small checks')
    add_from(exit_slide, tpl['strip'], 'routes', 0.729, 1.90, 12.0, 0.36,
             text='CHOOSE ONE ROUTE, ONCE     /     SUPPORTED   ·   STANDARD   ·   STRETCH     /     YOU MAY PASS', sz=1200, bold=True)
    y = 2.36
    for route in ROUTES:
        data = KO['routes'][route]
        for i, question in enumerate(data['questions']):
            line = question
            if data.get('options'):
                line += '   ' + ' / '.join(data['options'][i])
            add_from(exit_slide, tpl['cue'], f'cue-{route}-{i}', 0.729, y, 1.20, 0.40, text=f'{data["label"][:3].upper()} {i + 1}', sz=1200)
            add_from(exit_slide, tpl['text'], f'line-{route}-{i}', 2.0, y, 10.65, 0.50, text=line, sz=1500)
            y += 0.56
    add_from(exit_slide, tpl['line'], 'access', 0.729, 5.80, 11.9, 0.95, text=KO['exit_intro'] + ' Next lesson: ' + KO['next_lesson'] + '.', sz=1200)
    replace_notes(exit_slide, {
        'Exit answer: Random particle movement and crossings continue at equilibrium.':
        KO['staff'] + ' Pupil worksheet page 6 carries the three exit routes.\n[Staff explanations — do not display]\n'
        + '\n'.join(f'{KO["routes"][r]["label"]} {i + 1}. {a}' for r in ROUTES for i, a in enumerate(KO['routes'][r]['answers']))
        + '\nNext lesson: ' + KO['next_lesson'] + '.'})

    # ---- the retrieval check answered the retired arrival questions; it keeps the prediction
    set_text(title(check), 'Prediction check')
    remove_shapes(check, [s.name for s in check.shapes if s.has_text_frame and s.top in (1885950, 2647950)])
    for s in check.shapes:
        if s.top == 4429125: s.top = int(1.95 * EMU)
        elif s.top == 4619625: s.top = int(2.16 * EMU)
        elif s.name == 'goto:1' and s.top == 5581650: s.top = int(3.20 * EMU)
    replace_notes(check, {'Read only after an attempt; retain the pupil’s first view.':
                          'Read only after an attempt; retain the pupil’s first view. The arrival answers are in the notes of slides 1–3 and in the teacher guide.'})

    # ---- remaining day wording in retained notes, and the revision line beside the sources
    for slide in prs.slides:
        text = notes_text(slide)
        if DAYS.search(text):
            text = text.replace('Monday / Wednesday / Friday P1, 09:30–10:10', PLAN)
            set_notes(slide, text.split('\n'))
        if '[Sources]' in notes_text(slide):
            replace_notes(slide, {'[Sources]\n': REVISION + '\n[Sources]\n'}) if notes_text(slide).count('[Sources]\n') == 1 else None
    first = list(prs.slides)[0]
    if REVISION not in notes_text(first):
        set_notes(first, [*notes_text(first).split('\n'), REVISION])
    check_no_day_words(prs)
    assert len(prs.slides) == 21
    return prs


# ----------------------------------------------------------------------------- worksheet
def build_sheet(root):
    doc = Document(io.BytesIO(baseline_bytes(root, SHEET)))
    subtitle = doc.paragraphs[1]
    assert subtitle.text == 'Monday | 09:30–10:10 | Introduce | Editable pupil sheet', subtitle.text
    for run in subtitle.runs[1:]:
        run._r.getparent().remove(run._r)
    subtitle.runs[0].text = 'Lesson 1 of Week 4 | 40 minutes | Introduce | Editable pupil sheet'
    table_style = doc.tables[0].style

    def page_break():
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    page_break()
    doc.add_paragraph('W4L1 | Arrival — choose one route', style='Title')
    doc.add_paragraph('Four quick questions | Point, say, sign, draw, type or write | You may pass', style='Subtitle')
    doc.add_paragraph(ARRIVAL['response'] + ' ' + ARRIVAL['access'], style='Caption')
    for route in ROUTES:
        data = ARRIVAL['routes'][route]
        doc.add_paragraph(data['label'] + ' route', style='Heading 2')
        for i, question in enumerate(data['questions']):
            doc.add_paragraph(f'{i + 1}. {question}', style='Normal')
            doc.add_paragraph('______________________________________________________________', style='Normal')

    page_break()
    doc.add_paragraph('W4L1 | Knowledge organiser', style='Title')
    doc.add_paragraph('Keep this page; use it in any stage', style='Subtitle')
    doc.add_paragraph('Objective: ' + KO['objective'], style='Normal')
    doc.add_paragraph('Success criteria', style='Heading 2')
    for i, item in enumerate(KO['success']):
        doc.add_paragraph(f'{i + 1}. {item}', style='Normal')
    doc.add_paragraph('Key words', style='Heading 2')
    table = doc.add_table(rows=1, cols=3)
    table.style = table_style
    for cell, head in zip(table.rows[0].cells, ('Word', 'Tier', 'Meaning')):
        cell.text = head
    for word, tier, meaning in KO['vocabulary']:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = word, tier, meaning
    doc.add_paragraph('Key ideas', style='Heading 2')
    for fact in KO['facts']:
        doc.add_paragraph(fact, style='Normal')
    doc.add_paragraph(KO['model_caption'], style='Caption')
    doc.add_paragraph('Diffusion or not?', style='Heading 2')
    for example in KO['examples']:
        doc.add_paragraph(example, style='Normal')

    page_break()
    doc.add_paragraph('W4L1 | Exit ticket — complete one route, once', style='Title')
    doc.add_paragraph('Two small checks | You may pass', style='Subtitle')
    doc.add_paragraph(KO['exit_intro'], style='Caption')
    for route in ROUTES:
        data = KO['routes'][route]
        doc.add_paragraph(data['label'] + ' route', style='Heading 2')
        for i, question in enumerate(data['questions']):
            line = question
            if data.get('options'):
                line += '   (' + ' / '.join(data['options'][i]) + ')'
            doc.add_paragraph(f'{i + 1}. {line}', style='Normal')
            doc.add_paragraph('______________________________________________________________', style='Normal')
    doc.add_paragraph('Next lesson: ' + KO['next_lesson'] + '.', style='Normal')
    doc.add_paragraph('Verbal, written, drawn, signed, symbol or choice selected, demonstrated or re-attempted, adult scribed: all count.', style='Caption')
    return doc


# ----------------------------------------------------------------------------- teacher guide
def build_guide(root):
    doc = Document(io.BytesIO(baseline_bytes(root, GUIDE)))
    paragraphs = doc.paragraphs
    start = next(i for i, p in enumerate(paragraphs) if p.text == 'W4L1 | Diffusion' and p.style.name == 'Title')
    end = next(i for i, p in enumerate(paragraphs) if i > start and p.style.name == 'Title')
    anchor = paragraphs[end]   # 'W4L2 | Gas Exchange'
    assert anchor.text.startswith('W4L2'), anchor.text
    table_style = doc.tables[0].style if doc.tables else None

    def para(text, style):
        return anchor.insert_paragraph_before(text, style=style)

    para('Arrival answers (staff only; do not display)', 'Heading 2')
    for route in ROUTES:
        data = ARRIVAL['routes'][route]
        for i, answer in enumerate(data['answers']):
            para(f'{data["label"]} {i + 1}. {answer}', 'Normal')
    para(ARRIVAL['staff'], 'Caption')
    para('Exit ticket answers (staff only)', 'Heading 2')
    for route in ROUTES:
        data = KO['routes'][route]
        for i, answer in enumerate(data['answers']):
            para(f'{data["label"]} {i + 1}. {answer}', 'Normal')
    para(KO['staff'] + ' Next lesson: ' + KO['next_lesson'] + '.', 'Caption')
    para(FB['staff_title'], 'Heading 2')
    para(FB['policy_line'], 'Caption')
    table = doc.add_table(rows=1, cols=2)
    if table_style is not None:
        table.style = table_style
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = 'Code', 'Meaning'
    for code, meaning in FB['codes']:
        row = table.add_row().cells
        row[0].text, row[1].text = code, meaning
    anchor._p.addprevious(table._tbl)
    para(FB['summary_line'], 'Normal')
    para(FB['launch_line'], 'Normal')
    for heading, text in FB['staff_sections']:
        para(f'{heading} — {text}', 'Normal')
    para('Revision 15 September 2026: this section is aligned with the canonical lesson’s arrival routes, knowledge '
         'organiser, exit routes and feedback card. Templates carry no pupil data.', 'Caption')
    return doc


# ----------------------------------------------------------------------------- render
def render_pdf(source, output):
    with tempfile.TemporaryDirectory(prefix='w4l1-render-') as temp:
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', temp, str(source)],
                       check=True, capture_output=True, timeout=300)
        produced = Path(temp) / (source.stem + '.pdf')
        assert produced.is_file(), produced
        shutil.copyfile(produced, output)


def combine_worksheets(root):
    import pymupdf
    pack = root / PACK
    order = [f'Week_{w}/W{w}L{p}_Worksheet.pdf' for w in range(3, 8) for p in range(1, 4)]
    out = pymupdf.open()
    for rel in order:
        with pymupdf.open(pack / rel) as part:
            out.insert_pdf(part)
    out.save(pack / 'Pupil_Worksheets.pdf', garbage=3, deflate=True)
    return out.page_count


if __name__ == '__main__':
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--root', type=Path, default=HERE.parents[1])
    args = ap.parse_args()
    root = args.root.resolve()
    deck = root / DECK
    buf = io.BytesIO(); build_deck(root).save(buf); deck.write_bytes(buf.getvalue())
    render_pdf(deck, deck.with_suffix('.pdf'))
    sheet = root / SHEET
    build_sheet(root).save(sheet)
    render_pdf(sheet, sheet.with_suffix('.pdf'))
    guide = root / GUIDE
    build_guide(root).save(guide)
    render_pdf(guide, guide.with_suffix('.pdf'))
    pages = combine_worksheets(root)
    print(json.dumps({'deck': DECK, 'sheet': SHEET, 'guide': GUIDE, 'pupilWorksheetPages': pages}, indent=2))
