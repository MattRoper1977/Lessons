#!/usr/bin/env python3
"""Gate for the companion packs placed by UX2 lane D2.

Reads ``data/companion-packs.json`` (the placement manifest) and asserts, on
the shipped tree:

  C1 byte witness   every listed file exists with the listed bytes and sha256
  C2 opens          every .pptx/.docx opens; every .pdf has pages > 0
  C3 forty minutes  every lesson .pptx: footers matching
                    ``^(?P<stage>.+?) · (?P<n>\\d+) min across this stage$``,
                    deduplicated by stage, sum to 40
  C4 learner names  (s23) no docx/pdf carries a filled Name/Pupil/Learner/Student
                    field; blanks, underscores and dots are fine
  C5 forbidden      OUTSTANDING_V4 / _V3_1 / Download.rar appear 0 times in pptx
                    slide bodies, docx paragraphs+tables, pdf text (speaker notes
                    are counted and printed separately, never failed on)
  C6 holds          every held file is absent, and no unlisted native file sits
                    in a pack directory
  C7 manifest       schema 1; every pack has >= 1 file; every file belongs to
                    exactly one pack; ids unique; companionOf exists

Weeks come from the manifest's ``wtoken`` and are never parsed from file
names (g27). Exit 0 PASS, 1 FAIL, 2 INCONCLUSIVE (dependency or manifest
missing). ``--self-test`` copies the listed files to a temporary root, plants
one defect per control and proves each control goes red.

Dependencies: python-pptx, python-docx, pypdf.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
MANIFEST = 'data/companion-packs.json'
FORBID = ('OUTSTANDING_V4', '_V3_1', 'Download.rar')
FOOTER_RE = re.compile(r'^(?P<stage>.+?) · (?P<n>\d+) min across this stage$')
NAME_RE = re.compile(r'(Name|Pupil|Learner|Student)\s*[:：]\s*(\S.*)')
NEXT_FIELD = re.compile(r'(Date|Class|Group|Form|Teacher|Tutor|Year)\s*[:：].*$')
BLANKISH = re.compile(r'[_\s.\-–—…☐□▢·:]+')
SUFFIX = {('lesson', 'pptx'): '.pptx', ('slides', 'pdf'): '.pdf', ('teacher', 'docx'): '_Teacher.docx',
          ('teacher', 'pdf'): '_Teacher.pdf', ('pupil', 'docx'): '_Pupil.docx', ('pupil', 'pdf'): '_Pupil.pdf'}
NATIVE = {'.pptx', '.docx', '.pdf'}


def load_deps():
    try:
        from pptx import Presentation  # noqa: F401
        import docx  # noqa: F401
        from pypdf import PdfReader  # noqa: F401
    except ImportError as e:
        print(f'INCONCLUSIVE: missing dependency {e.name!r} (need python-pptx, python-docx, pypdf)')
        sys.exit(2)


def sha256(p: Path) -> str:
    h = hashlib.sha256()
    with p.open('rb') as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()


def pptx_text(p: Path):
    """(slide count, body lines, notes lines)."""
    from pptx import Presentation
    prs = Presentation(str(p))
    body, notes = [], []
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                body.append(sh.text_frame.text)
            if getattr(sh, 'has_table', False) and sh.has_table:
                body += [c.text for r in sh.table.rows for c in r.cells if c.text.strip()]
        if sl.has_notes_slide and sl.notes_slide.notes_text_frame is not None:
            t = sl.notes_slide.notes_text_frame.text
            if t.strip():
                notes.append(t)
    return len(prs.slides), body, notes


def docx_lines(p: Path) -> list[str]:
    import docx
    d = docx.Document(str(p))
    out = [x.text for x in d.paragraphs]
    out += [c.text for t in d.tables for r in t.rows for c in r.cells]
    for sec in d.sections:
        for part in (sec.header, sec.footer):
            out += [x.text for x in part.paragraphs]
    return out


def pdf_lines(p: Path):
    from pypdf import PdfReader
    r = PdfReader(str(p))
    out = []
    for pg in r.pages:
        out.extend((pg.extract_text() or '').split('\n'))
    return len(r.pages), out


def gate40(body: list[str]):
    stages, conflicts = {}, []
    for t in body:
        m = FOOTER_RE.match(t.strip())
        if m:
            stage, n = m.group('stage'), int(m.group('n'))
            if stage in stages and stages[stage] != n:
                conflicts.append(f'{stage}: {stages[stage]} vs {n}')
            stages.setdefault(stage, n)
    return sum(stages.values()), stages, conflicts


def s23_hits(lines: list[str]) -> list[str]:
    hits = []
    for ln in lines:
        for m in NAME_RE.finditer(ln):
            val = NEXT_FIELD.sub('', m.group(2))
            if BLANKISH.sub('', val):
                hits.append(f'{m.group(1)}: {val.strip()[:60]}')
    return hits


def forbidden_hits(lines: list[str]) -> int:
    return sum(1 for ln in lines for tok in FORBID if tok in ln)


def run(root: Path, manifest_path: Path, quiet: bool = False) -> dict:
    """Run every control; return {'fail': {control: [messages]}, 'stats': {...}}."""
    fail = {k: [] for k in ('C1', 'C2', 'C3', 'C4', 'C5', 'C6', 'C7')}
    stats = {'packs': 0, 'files': 0, 'bytes': 0, 'pptx': 0, 'docx': 0, 'pdf': 0, 'pdfPages': 0, 'slides': 0,
             'footerSums': {}, 's23Labels': 0, 'forbiddenBody': 0, 'forbiddenNotes': 0, 'held': 0, 'wtokens': []}
    if not manifest_path.is_file():
        print(f'INCONCLUSIVE: manifest missing {manifest_path}')
        sys.exit(2)
    m = json.loads(manifest_path.read_text(encoding='utf-8'))
    if m.get('schema') != 1:
        fail['C7'].append(f'schema {m.get("schema")!r} != 1')
    packs = m.get('packs', [])
    stats['packs'] = len(packs)
    owners: dict[str, list[str]] = {}
    ids = [p.get('id') for p in packs]
    if len(set(ids)) != len(ids):
        fail['C7'].append('duplicate pack ids')
    dir_files: dict[Path, set[str]] = {}
    for p in packs:
        pid = p.get('id', '?')
        if not p.get('files'):
            fail['C7'].append(f'{pid}: no placed files')
        if p.get('term') not in ('Autumn 1', 'Autumn 2') or not p.get('wtoken'):
            fail['C7'].append(f'{pid}: term/wtoken missing')
        stats['wtokens'].append(p.get('wtoken'))
        live = root / p.get('companionOf', '')
        if not p.get('companionOf') or not live.is_file():
            fail['C7'].append(f'{pid}: companionOf missing in tree {p.get("companionOf")}')
        lesson_stem = None
        for f in p['files']:
            rel = f['path']
            if '..' in Path(rel).parts or Path(rel).is_absolute():
                fail['C7'].append(f'{pid}: unsafe path {rel}')
                continue
            owners.setdefault(rel, []).append(pid)
            path = root / rel
            dir_files.setdefault(path.parent, set()).add(path.name)
            if (f['role'], f['type']) not in SUFFIX:
                fail['C7'].append(f'{pid}: unknown role/type {f["role"]}/{f["type"]}')
            if f['role'] == 'lesson':
                lesson_stem = path.with_suffix('').name
            stats['files'] += 1
            # C1
            if not path.is_file():
                fail['C1'].append(f'missing {rel}')
                continue
            size = path.stat().st_size
            digest = sha256(path)
            stats['bytes'] += size
            if size != f['bytes'] or digest != f['sha256']:
                fail['C1'].append(f'byte witness {rel}: {size} B {digest[:12]} vs manifest {f["bytes"]} B {f["sha256"][:12]}')
                continue
            # C2..C5 by type
            try:
                if path.suffix == '.pptx':
                    stats['pptx'] += 1
                    n, body, notes = pptx_text(path)
                    stats['slides'] += n
                    if f['role'] == 'lesson':
                        total, stages, conflicts = gate40(body)
                        stats['footerSums'][total] = stats['footerSums'].get(total, 0) + 1
                        if total != 40 or conflicts:
                            fail['C3'].append(f'{rel}: stage minutes sum {total} {stages} conflicts={conflicts}')
                    nb = forbidden_hits(body)
                    stats['forbiddenBody'] += nb
                    if nb:
                        fail['C5'].append(f'{rel}: {nb} forbidden token(s) in slide body')
                    stats['forbiddenNotes'] += forbidden_hits(notes)
                elif path.suffix == '.docx':
                    stats['docx'] += 1
                    lines = docx_lines(path)
                    stats['s23Labels'] += sum(1 for ln in lines for _ in NAME_RE.finditer(ln))
                    hits = s23_hits(lines)
                    if hits:
                        fail['C4'].append(f'{rel}: filled learner name {hits[:3]}')
                    nb = forbidden_hits(lines)
                    stats['forbiddenBody'] += nb
                    if nb:
                        fail['C5'].append(f'{rel}: {nb} forbidden token(s) in document text')
                elif path.suffix == '.pdf':
                    stats['pdf'] += 1
                    n, lines = pdf_lines(path)
                    stats['pdfPages'] += n
                    if n <= 0:
                        fail['C2'].append(f'{rel}: 0 pages')
                    stats['s23Labels'] += sum(1 for ln in lines for _ in NAME_RE.finditer(ln))
                    hits = s23_hits(lines)
                    if hits:
                        fail['C4'].append(f'{rel}: filled learner name {hits[:3]}')
                    nb = forbidden_hits(lines)
                    stats['forbiddenBody'] += nb
                    if nb:
                        fail['C5'].append(f'{rel}: {nb} forbidden token(s) in pdf text')
                else:
                    fail['C7'].append(f'{pid}: unexpected suffix {rel}')
            except Exception as e:  # noqa: BLE001
                fail['C2'].append(f'{rel}: does not open ({type(e).__name__}: {e})')
        # C6 held absent
        for h in p.get('held', []):
            stats['held'] += 1
            if lesson_stem is None:
                fail['C6'].append(f'{pid}: cannot derive held path (no lesson pptx placed)')
                continue
            suffix = SUFFIX.get((h['role'], h['type']))
            if suffix is None:
                fail['C7'].append(f'{pid}: held entry with unknown role/type {h}')
                continue
            would_be = (root / p['files'][0]['path']).parent / (lesson_stem + suffix)
            if would_be.exists():
                fail['C6'].append(f'held file present {would_be.relative_to(root).as_posix()} ({h["reason"]})')
    for rel, who in owners.items():
        if len(who) != 1:
            fail['C7'].append(f'{rel} belongs to {len(who)} packs {who}')
    for d, names in dir_files.items():
        if d.is_dir():
            strays = sorted(x.name for x in d.iterdir() if x.is_file() and x.suffix in NATIVE and x.name not in names)
            if strays:
                fail['C6'].append(f'unlisted native file(s) in {d.relative_to(root).as_posix()}: {strays}')
    stats['wtokens'] = sorted(set(stats['wtokens']))
    if not quiet:
        print(json.dumps({'root': str(root), 'manifest': str(manifest_path.relative_to(root)), 'stats': stats,
                          'failures': {k: v for k, v in fail.items() if v}}, indent=1, ensure_ascii=False))
    return {'fail': fail, 'stats': stats}


def self_test(root: Path, manifest_path: Path) -> int:
    from pptx import Presentation
    import docx
    m = json.loads(manifest_path.read_text(encoding='utf-8'))
    tmp = Path(tempfile.mkdtemp(prefix='companion-packs-selftest-'))
    try:
        # mirror the listed files and the live lessons they accompany
        by_role = {}
        for p in m['packs']:
            src = root / p['companionOf']
            dst = tmp / p['companionOf']
            dst.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(src, dst)
            for f in p['files']:
                dst = tmp / f['path']
                dst.parent.mkdir(parents=True, exist_ok=True)
                shutil.copyfile(root / f['path'], dst)
                by_role.setdefault((f['role'], f['type']), []).append((p, f))
        plants = {}

        def restamp(f, path):
            f['bytes'] = path.stat().st_size
            f['sha256'] = path.stat().st_size and sha256(path)

        # C1: flip one byte in a placed pupil pdf; manifest untouched
        p, f = by_role[('pupil', 'pdf')][0]
        path = tmp / f['path']
        b = bytearray(path.read_bytes())
        b[len(b) // 2] ^= 0xFF
        path.write_bytes(bytes(b))
        plants['C1'] = f['path']
        # C3: +1 minute on the first footer of a lesson pptx; manifest restamped so only C3 sees it
        p, f = by_role[('lesson', 'pptx')][1]
        path = tmp / f['path']
        prs = Presentation(str(path))
        altered = None
        for sl in prs.slides:
            for sh in sl.shapes:
                if altered or not sh.has_text_frame:
                    continue
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        mm = FOOTER_RE.match(r.text.strip())
                        if mm and not altered:
                            r.text = f'{mm.group("stage")} · {int(mm.group("n")) + 1} min across this stage'
                            altered = r.text
        assert altered, 'self-test could not find a footer to alter'
        prs.save(str(path))
        restamp(f, path)
        plants['C3'] = f['path']
        # C4: fill a learner name in a pupil docx
        p, f = by_role[('pupil', 'docx')][2]
        path = tmp / f['path']
        d = docx.Document(str(path))
        d.add_paragraph('Name: Jamie Smith   Date: ______')
        d.save(str(path))
        restamp(f, path)
        plants['C4'] = f['path']
        # C5: forbidden token in a teacher docx body
        p, f = by_role[('teacher', 'docx')][3]
        path = tmp / f['path']
        d = docx.Document(str(path))
        d.add_paragraph('Primary classroom source: EXAMPLE_OUTSTANDING_V4.html')
        d.save(str(path))
        restamp(f, path)
        plants['C5'] = f['path']
        # C6: a held file appears at its would-be path
        held_pack = next(p for p in m['packs'] if p['held'])
        h = held_pack['held'][0]
        lesson = next(f for f in held_pack['files'] if f['role'] == 'lesson')
        would_be = (tmp / lesson['path']).parent / (Path(lesson['path']).stem + SUFFIX[(h['role'], h['type'])])
        would_be.write_bytes(b'%PDF-1.4\n%planted held file\n%%EOF\n')
        plants['C6'] = would_be.relative_to(tmp).as_posix()
        # C2: a listed pdf that is not a pdf; restamped so only C2 sees it
        p, f = by_role[('teacher', 'pdf')][4]
        path = tmp / f['path']
        path.write_bytes(b'not a pdf at all')
        restamp(f, path)
        plants['C2'] = f['path']
        # C7: the same path listed under two packs
        dup = m['packs'][5]['files'][0]
        m['packs'][6]['files'].append(dict(dup))
        plants['C7'] = dup['path']

        tmp_manifest = tmp / MANIFEST
        tmp_manifest.parent.mkdir(parents=True, exist_ok=True)
        tmp_manifest.write_text(json.dumps(m), encoding='utf-8')
        res = run(tmp, tmp_manifest, quiet=True)
        ok = True
        for ctrl in ('C1', 'C2', 'C3', 'C4', 'C5', 'C6', 'C7'):
            msgs = res['fail'][ctrl]
            hit = any(plants[ctrl] in msg for msg in msgs)
            print(f'{"PASS" if hit else "FAIL"} {ctrl} planted={plants[ctrl]} -> {len(msgs)} failure(s)'
                  + (f': {msgs[0][:160]}' if msgs else ''))
            ok &= hit
        print('self-test', 'PASS' if ok else 'FAIL', '(every planted defect went red)' if ok else '')
        return 0 if ok else 1
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--root', type=Path, default=ROOT)
    ap.add_argument('--manifest', type=Path, default=None)
    ap.add_argument('--self-test', action='store_true')
    a = ap.parse_args()
    load_deps()
    root = a.root.resolve()
    manifest = (a.manifest or root / MANIFEST).resolve()
    if a.self_test:
        return self_test(root, manifest)
    res = run(root, manifest)
    failures = sum(len(v) for v in res['fail'].values())
    s = res['stats']
    print(f'weeks read from manifest (never from file names): {s["wtokens"]}')
    print(f'speaker-note forbidden hits (reported, not failed): {s["forbiddenNotes"]}')
    print('PASS' if failures == 0 else f'FAIL ({failures} failure(s))',
          f'- {s["packs"]} packs, {s["files"]} files, {s["bytes"]} B, {s["pptx"]} pptx/{s["slides"]} slides, '
          f'{s["docx"]} docx, {s["pdf"]} pdf/{s["pdfPages"]} pages, {s["held"]} held')
    return 0 if failures == 0 else 1


if __name__ == '__main__':
    sys.exit(main())
