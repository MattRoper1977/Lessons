"""Bounded paired-lesson feedback authoring. Rebase BASE before use over later edits."""
from pathlib import Path
import io,json,subprocess,re,html
from docx import Document

BASE='c7e64b2baf6a3f17f261fb001d416886d46fd32a'
HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
REL='Science_Teesside/Grow/SCI_G_W3_Friction.html'
PACK='Science_Teesside/Teaching_Packs/GROW/lessons/'
D=json.loads((HERE/'W3_FEEDBACK.json').read_text())
e=html.escape
def baseline(path):return subprocess.check_output(['git','-C',str(ROOT),'show',BASE+':'+path])
def prompt(which):return '<details class="grow-help-prompt" id="grow-help-'+which+'"><summary>'+e(D['prompt_title'])+'</summary><p>'+e(D['prompt_'+which])+'</p></details>'
def staff():return ''.join('<h3>'+e(h)+'</h3><p>'+e(t)+'</p>' for h,t in D['staff_sections'])
def review():return '<p>'+e(D['review_intro'])+'</p><ol>'+''.join('<li>'+e(t)+'</li>' for t in D['review_steps'])+'</ol><p>'+e(D['review_close'])+'</p>'
s=baseline(REL).decode()
# Replace the participation-only stage with science review without changing indices or timing.
a=s.index('<div class="slide" data-title="Lundy Loop"');b=s.index('<div class="slide" data-title="Exit Ticket"',a)
s=s[:a]+'<div class="slide" data-title="Review the evidence" data-timer="4" id="grow-review-evidence"><span class="slide-tag tag-wedo">Review · use your results</span><h2>'+e(D['review_title'])+'</h2>'+review()+prompt('b')+'</div>'+s[b:]
s=s.replace('<details id="grow-w3a-exit"',prompt('a')+'<details id="grow-w3a-exit"',1)
# Teacher guidance is explicit and optional, outside the default pupil pack.
teacher='<details class="grow-feedback-staff" id="grow-feedback-staff"><summary>Staff: feedback and marking</summary>'+staff()+'<button type="button" class="ghost small" onclick="printSection(\'teacher-feedback\',\'standard\')">Print staff feedback guidance</button><p>Optional existing staff records; use only when appropriate.</p><button type="button" class="ghost small" onclick="printSection(\'witness\',\'standard\')">Print assessor witness record</button> <button type="button" class="ghost small" onclick="printSection(\'feedback\',\'standard\')">Print optional marking notes</button></details>'
s=s.replace('<nav class="sr-offline-nav"',teacher+'<nav class="sr-offline-nav"',1)
a=s.index('<div id="print-lundy"');b=s.index('<div id="print-feedback"',a)
s=s[:a]+'<div id="print-review-evidence" class="print-section"><h2>'+e(D['review_title'])+'</h2>'+review()+'<p><strong>Optional:</strong> '+e(D['prompt_b'])+'</p></div><div id="print-teacher-feedback" class="print-section"><h2>'+e(D['staff_title'])+'</h2>'+staff()+'</div>'+s[b:]
s=s.replace("['ko','intro','glance','arrival','wedo','exit','witness','lundy','feedback']", "['ko','intro','glance','arrival','wedo','w3a-exit','review-evidence','exit']",1)
s=s.replace('— Feedback Sheet</h2><p style="font-size:.95rem;color:#444">Keep with your portfolio.</p>', '— Optional staff marking notes</h2><p style="font-size:.95rem;color:#444">Staff use only when appropriate. Use existing work or a spoken, signed, pointed or drawn response; no duplicate pupil form is required.</p>',1)
s=s.replace('Pupil response — my next step:', 'Response observed — record only if useful:',1)
# Remove obsolete participation-only styling now that no element uses it.
s=re.sub(r'\.tag-lundy\{[^}]*\}\s*\.lundy-grid\{[^}]*\}\s*\.lundy-box\{[^}]*\}\s*\.lundy-box h3\{[^}]*\}\s*\.lundy-box p\{[^}]*\}', '',s,count=1)
s=s.replace('@media print{.lundy-grid{display:block}}','')
replacements={
 'Monday resources':'Lesson A resources','Tuesday resources':'Lesson B resources',
 'Monday P2 (10:10–10:50); then Tuesday P1 (09:30–10:10). Keep this deck and print pack for both days.':'Lesson A: Friction: Friend and Enemy. Lesson B: Friction: test the surfaces. Keep this deck and your work for both lessons.',
 'Resume period 2 — Tuesday P1 (09:30–10:10)':'Resume Lesson B — Friction: test the surfaces',
 '▶ Tuesday P1 (09:30–10:10) — restart here.':'Lesson B: Friction: test the surfaces — restart here.',
 'Retrieve Monday’s prediction.':'Retrieve your Lesson A prediction.',
 'Monday P2 (10:10–10:50) — Title through We Do 2: 40 minutes. STOP after We Do 2, label and keep the work. Tuesday P1 (09:30–10:10) — restart at Independent Work: 32-minute workshop, 4-minute pupil feedback, 4-minute Exit Ticket. Total: 40 minutes on each day; 80 across the week.':'Lesson A — Title through We Do 2: 40 minutes, including its final two-minute exit. STOP after We Do 2 and keep your work. Lesson B: Friction: test the surfaces — restart at Independent Work: 32-minute workshop, 4-minute evidence review, 4-minute Exit Ticket. Total: 40 minutes in each lesson; 80 across the pair.',
 '1 minutes. Two separate lessons: Monday P2 (10:10–10:50) and Tuesday P1 (09:30–10:10).':'1 minute. Two separate 40-minute lessons: Friction: Friend and Enemy; then Friction: test the surfaces.',
 'label/save work for Tuesday P1 (09:30–10:10).':'keep work for Friction: test the surfaces.',
 'BREAK POINT: Period 1 ends after this slide.':'BREAK POINT: Lesson A ends after this slide.',
 '"Lundy Loop": "4 minutes. Receive one actual pupil observation or access request. Name and enact a feasible response; explain any limit. Do not assume a display, outside audience or compelled peer discussion."':'"Review the evidence": "4 minutes. Revisit one conclusion against the existing results or clearly labelled model data. Hear, read or watch the pupil response; adapt the explanation or scaffold and explain any limit. Preserve original values and label retries. Passing is allowed. No separate feedback form or automatic R. Then use the existing four-minute exit."'
}
for old,new in replacements.items():
 assert old in s,old
 s=s.replace(old,new)
# Named involvement is part of ordinary W3A guided practice, not added time.
s=s.replace('A wrong sort is data, not a fail. BREAK POINT:', 'A wrong sort is data, not a fail. Invite the pupil to identify one word or idea to check; engage with the response and adapt the explanation or scaffold. The help prompt is optional, passing is allowed, and no extra written record is required. BREAK POINT:',1)
css='''<style id="grow-feedback-css">
.grow-help-prompt,.grow-feedback-staff{margin:14px 0;padding:10px 12px;border:1px solid #778891;border-radius:8px}
.grow-help-prompt summary,.grow-feedback-staff summary{font-weight:700;cursor:pointer;min-height:44px;padding:10px 0}
.grow-help-prompt summary:focus-visible,.grow-feedback-staff summary:focus-visible{outline:3px solid var(--btn-bg);outline-offset:2px}
.grow-help-prompt p,.grow-feedback-staff p{font-size:1rem;line-height:1.45}.grow-feedback-staff h3{font-size:1rem}
#grow-review-evidence li{margin-bottom:12px}
@media print{#print-teacher-feedback p,#print-teacher-feedback h3{font-size:10pt!important;line-height:1.3!important;margin:6px 0!important}#print-review-evidence p,#print-review-evidence li{font-size:11pt!important;line-height:1.4!important}}
</style>'''
s=s.replace('</head>',css+'</head>')
(ROOT/REL).write_text(s)

paths={
 'a_pupil':PACK+'W3A/GROW_W3A_Friction_Pupil.docx',
 'a_teacher':PACK+'W3A/GROW_W3A_Friction_Teacher.docx',
 'b_pupil':PACK+'W3B/GROW_Science_Autumn1_W3B_Friction_Test_Pupil.docx',
 'b_teacher':PACK+'W3B/GROW_Science_Autumn1_W3B_Friction_Test_Teacher.docx'
}
for key,path in paths.items():
 doc=Document(io.BytesIO(baseline(path)))
 def replace(i,text):doc.paragraphs[i].text=text
 if key=='a_pupil':
  replace(2,'Week 3 · Lesson A · 40 minutes')
  replace(24,'Keep the decision tied to the contact and the job. No oil is needed. Optional: '+D['prompt_a'])
  for i in [48,71,93]:replace(i,'Finish one two-check exit route, then stop and keep your work. Next lesson: Friction: test the surfaces.')
 elif key=='a_teacher':
  replace(3,'Teach Lesson A’s explanation and shared reasoning. Stop after We Do 2 and its exit. Lesson B’s 32-minute workshop, 4-minute evidence review and 4-minute exit are separate.')
  replace(5,'Objective: explain what friction does, with one helpful and one unhelpful example. Success includes identifying contacts, predicting a surface comparison and supporting a claim. An adult demonstration is not evidence that a pupil performed the test.')
  replace(45,'Close Lesson A; retain the evidence')
  replace(46,'Feedback revision · 15 September 2026 · draft pilot')
  replace(48,'Stop after We Do 2 and its exit. Keep the prediction for Friction: test the surfaces: workshop 32 minutes, evidence review 4 minutes, exit 4 minutes. The investigation has its own Lesson B pack.')
  replace(53,'Canonical lesson: Science_Teesside/Grow/SCI_G_W3_Friction.html. Lesson A companion: resources/GS_W3A.html. Historical conversion source: 5018839965d424afd31bc0ac1b74dc8f713ceca2.')
  replace(55,'The seven Lesson A timers and five helpful/unhelpful situations are retained. Tasks map to these stages; a chosen response page replaces lengthy copying. Lesson B’s results and exit stay in its own pack. Feedback guidance is on the final page.')
 elif key=='b_pupil':
  replace(5,"Use this sheet with your Lesson A prediction. You can test, observe, direct an adult, point or dictate. Keep real results even if they surprise you. Use model data only if a safe practical is unavailable.")
  replace(7,"Retrieve your Lesson A prediction. Which surface might stop the block sooner? Explain or point to a reason. Name what must stay the same.")
  replace(55,'5. Review the evidence and exit')
  replace(57,'Use four minutes to review one conclusion against your existing results with the adult. Keep or improve the explanation, preserving original values. Then use four minutes for one exit question.')
  replace(58,'Optional: '+D['prompt_b'])
  replace(78,'Keep prediction, results and conclusion together. Next lesson: Levers, pulleys and gears.')
 else:
  replace(13,'Lesson A prediction and pupil record; optional offline XLSX for actual measurements')
  replace(19,'Review the evidence · 4 minutes')
  replace(20,'Revisit one conclusion against the existing results or labelled model data. Engage with the response and adapt the explanation or scaffold. Preserve original values; no separate loop sheet is needed.')
  replace(36,'5. Review the evidence and exit')
  replace(37,'Review: check the conclusion against both trials and the controls; retain uncertainty honestly. Exit: brakes/shoes/tyres/handles; brake contact resists sliding and slows the wheel, with tyre-road grip also needed; contrast this with unwanted engine heating/wear. Accept equivalent responses.')
  replace(56,'Finish at 40 minutes: 32 workshop + 4 evidence review + 4 exit. Next lesson: Levers, pulleys and gears.')
  replace(60,'Historical conversion source Git blob: 170d772a5e9d23909e95a534c5c1e16ee7df67c8')
  replace(62,'Feedback revision: 2026-09-15 · draft pilot')
  replace(64,'Workshop remains 32 minutes, followed by four-minute evidence review and four-minute exit. The final page gives integrated feedback guidance; no separate participation stage or loop record is required.')
  doc.tables[0].rows[2].cells[1].text='Review the evidence'
 if key.endswith('teacher'):
  p=doc.add_paragraph('STAFF GUIDANCE / INTEGRATED FEEDBACK','Kicker');p.paragraph_format.page_break_before=True
  doc.add_paragraph(D['staff_title'],'Title')
  for h,t in D['staff_sections']:
   doc.add_paragraph(h,'Heading 2');doc.add_paragraph(t)
 doc.save(ROOT/path)
