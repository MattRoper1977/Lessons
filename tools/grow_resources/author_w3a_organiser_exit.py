"""EDU-Q1 bounded organiser/exit authoring. Input is the accepted arrival draft.
Review/rebase BASE before using over later edits. Native PDFs come from DOCX.
"""
from pathlib import Path
import argparse,subprocess,json,io,re,html
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import fitz
BASE='194075db287e091a1f1fc1020c5873a00cdb1447'
HERE=Path(__file__).resolve().parent
REL='Science_Teesside/Grow/SCI_G_W3_Friction.html'
PACK='Science_Teesside/Teaching_Packs/GROW/lessons/W3A/'

def main(root):
 d=json.loads((HERE/'W3A_ORGANISER_EXIT.json').read_text());e=html.escape
 def read(rel):return subprocess.check_output(['git','-C',str(root),'show',BASE+':'+rel])
 s=read(REL).decode();svg=(HERE/'friction_contact.svg').read_text()
 def ko(prefix):
  model=svg.replace('contact-title',prefix+'-model-title')
  return '<p><strong>Learning objective:</strong> '+e(d['objective'])+'</p><h3>Success criteria</h3><ul>'+''.join('<li>'+e(x)+'</li>' for x in d['success'])+'</ul><div class="grow-ko-table-wrap" role="region" aria-label="Friction vocabulary"><table class="ko-table"><caption>Key words</caption><thead><tr><th scope="col">Word</th><th scope="col">Meaning</th></tr></thead><tbody>'+''.join('<tr><th scope="row">'+e(w)+'</th><td>'+e(v)+'</td></tr>' for w,v in d['vocabulary'])+'</tbody></table></div><h3>Key ideas</h3><ul>'+''.join('<li>'+e(x)+'</li>' for x in d['facts'])+'</ul><figure class="grow-contact-model">'+model+'<figcaption>'+e(d['model'])+'</figcaption></figure><h3>Helpful or unhelpful for the job?</h3><ul>'+''.join('<li>'+e(x)+'</li>' for x in d['examples'])+'</ul><p><strong>Use evidence:</strong> '+e(d['comparison'])+'</p>'
 # Reach the organiser from the seven W3A stages without changing the chassis controls.
 start=0
 for title in ['Title','Arrival Task','Today at a Glance','I Do 1','We Do 1','I Do 2','We Do 2']:
  at=s.index('data-title="'+title+'"',start);h=re.search(r'</h[12]>',s[at:]);assert h
  end=at+h.end();s=s[:end]+'<div class="grow-ko-access"><button type="button" class="ghost small" onclick="openGrowKO(this)">Knowledge organiser</button></div>'+s[end:];start=end
 dialog='<dialog id="grow-ko-dialog" aria-labelledby="grow-ko-heading"><form method="dialog" class="grow-ko-close"><button class="ghost small">Close organiser</button></form><h2 id="grow-ko-heading" tabindex="-1">Friction knowledge organiser</h2>'+ko('screen-ko')+'<button type="button" class="ghost small" onclick="printGrowSheet(\'print-ko\',\'supported\')">Print organiser</button></dialog>'
 s=s.replace('<div id="print-area">',dialog+'<div id="print-area">',1)
 a=s.index('<div id="print-ko"');b=s.index('<div id="print-intro"',a)
 s=s[:a]+'<div id="print-ko" class="print-section grow-ko-print"><h2>Friction knowledge organiser</h2>'+ko('paper-ko')+'</div>'+s[b:]
 exit_html='<details id="grow-w3a-exit" class="grow-exit-panel"><summary>Finish Lesson A · two small checks</summary><h3>Friction exit</h3><p>'+e(d['exit_intro'])+'</p><div role="group" aria-label="Exit route" class="level-toggle">'+''.join(f'<button type="button" class="ghost small" data-grow-exit-route="{k}" aria-pressed="{str(k=="supported").lower()}" onclick="setGrowExit(\'{k}\')">{r["label"]}</button>' for k,r in d['routes'].items())+'</div>'
 for k,r in d['routes'].items():
  exit_html+=f'<div id="grow-exit-{k}"'+(' hidden' if k!='supported' else '')+'>'
  for i,q in enumerate(r['questions']):
   exit_html+=f'<fieldset><legend>{i+1}. {e(q)}</legend>'
   if k=='supported':
    for option in r['options'][i]:exit_html+=f'<label class="grow-exit-choice"><input type="radio" name="grow-exit-{k}-{i}" value="{e(option)}"> {e(option)}</label>'
   else:exit_html+=f'<label class="grow-exit-input-label" for="grow-exit-{k}-{i}">Your response</label><textarea id="grow-exit-{k}-{i}" name="grow-exit-{k}-{i}" rows="2"></textarea>'
   exit_html+='</fieldset>'
  exit_html+=f'<button type="button" class="ghost small" onclick="printGrowSheet(\'print-w3a-exit\',\'{k}\')">Print {r["label"]} exit</button> <button type="button" class="ghost small" onclick="saveGrowExit(\'{k}\')">Save {r["label"]} responses</button><details class="grow-exit-explanations"><summary>Compare {r["label"]} explanations</summary><ol>'+''.join('<li>'+e(a)+'</li>' for a in r['answers'])+'</ol><p>Use these with an adult. Your response is not automatically marked.</p><button type="button" class="ghost small" onclick="printGrowSheet(\'print-w3a-exit-answers\',\'supported\')">Print Lesson A staff exit answers</button></details></div>'
 exit_html+='<p id="grow-exit-save-status" role="status"></p><p>Next lesson: '+e(d['next_lesson'])+'</p></details>'
 a=s.index('<div class="slide" data-title="We Do 2"');b=s.index('<div class="slide" data-title="Independent Work"',a)
 segment=s[a:b];period=re.search(r'<div class="period-break".*?</div>',segment).group()
 segment=segment.replace(period,'');pos=segment.rfind('</div>')
 period='<div class="period-break" style="margin-top:14px;text-align:center;font-weight:800;color:#3F7D6E;border-top:2px dashed #3F7D6E;padding-top:8px">Stop after Lesson A. Keep your work. Next lesson: '+e(d['next_lesson'])+'.</div>'
 segment=segment[:pos]+exit_html+period+segment[pos:];s=s[:a]+segment+s[b:]
 def ticket(k,r):
  t='<section class="grow-exit-ticket"><h3>Friction exit · '+r['label']+' · Lesson A</h3><p>'+e(d['exit_intro'])+'</p>'
  for i,q in enumerate(r['questions']):t+='<p><strong>'+str(i+1)+'. '+e(q)+'</strong></p>'+('<p>'+(' / '.join(r['options'][i]))+'</p>' if k=='supported' else '<div class="grow-exit-writing-space" aria-hidden="true"></div>')
  return t+'<p>Next lesson: '+e(d['next_lesson'])+'</p></section>'
 prints='<div id="print-w3a-exit" class="print-section">'+''.join('<div class="'+k+'-content">'+ticket(k,r)*2+'</div>' for k,r in d['routes'].items())+'</div><div id="print-w3a-exit-answers" class="print-section"><h2>Friction Lesson A exit answers</h2><p>'+e(d['staff'])+'</p>'+''.join('<h3>'+r['label']+'</h3><ol>'+''.join('<li><strong>'+e(q)+'</strong><p>'+e(a)+'</p></li>' for q,a in zip(r['questions'],r['answers']))+'</ol>' for r in d['routes'].values())+'</div>'
 s=s.replace('<div id="print-wedo"',prints+'<div id="print-wedo"',1)
 # A real native modal blocks deck shortcuts even while its heading is focused.
 s=s.replace("if(document.querySelector('.v4-modal-overlay.visible')", "if(document.querySelector('dialog[open]')||document.querySelector('.v4-modal-overlay.visible')",1)
 js='''let growKOReturn=null;
function openGrowKO(button){
  const dialog=document.getElementById('grow-ko-dialog');
  if(dialog.open)return;
  growKOReturn=button;dialog.showModal();dialog.scrollTop=0;
  document.getElementById('grow-ko-heading').focus({preventScroll:true});
}
document.getElementById('grow-ko-dialog').addEventListener('close',()=>{if(growKOReturn&&growKOReturn.isConnected)growKOReturn.focus({preventScroll:true});});
function setGrowExit(route){
  ['supported','standard','stretch'].forEach(k=>{document.getElementById('grow-exit-'+k).hidden=k!==route;});
  document.querySelectorAll('[data-grow-exit-route]').forEach(b=>b.setAttribute('aria-pressed',String(b.dataset.growExitRoute===route)));
  document.querySelectorAll('.grow-exit-explanations').forEach(x=>x.open=false);
  document.getElementById('grow-exit-save-status').textContent='';
}
function printGrowSheet(id,route){
  if(!['print-ko','print-w3a-exit','print-w3a-exit-answers'].includes(id))return;
  if(!['supported','standard','stretch'].includes(route))return;
  document.body.classList.remove('print-supported','print-standard','print-stretch');document.body.classList.add('print-'+route);
  document.querySelectorAll('.print-section').forEach(x=>x.classList.remove('visible'));
  document.getElementById(id).classList.add('visible');window.print();
}
function saveGrowExit(route){
  if(!['supported','standard','stretch'].includes(route))return;
  const fields=[...document.querySelectorAll('#grow-exit-'+route+' fieldset')];
  const responses=fields.map(f=>{const input=f.querySelector('input:checked,textarea');return {question:f.querySelector('legend').textContent,response:input?input.value.trim():''};});
  const text='Friction — Lesson A — '+route+'\\nRecorded responses; no automatic mark.\\n\\n'+responses.map(r=>r.question+'\\n'+(r.response||'[No response recorded]')).join('\\n\\n');
  const url=URL.createObjectURL(new Blob([text],{type:'text/plain;charset=utf-8'}));const link=document.createElement('a');link.href=url;link.download='Friction_Lesson_A_'+route+'_responses.txt';link.click();setTimeout(()=>URL.revokeObjectURL(url),1000);
  document.getElementById('grow-exit-save-status').textContent='Responses prepared for download. No mark has been assigned.';
}
'''
 s=s.replace('function setArrivalLevel(level)',js+'function setArrivalLevel(level)',1)
 css='''<style id="grow-w3a-organiser-exit-css">
.grow-ko-access{margin:0 0 10px}.grow-ko-access button{min-height:44px}
#grow-ko-dialog{color:var(--text);background:var(--slide-bg);border:2px solid #3f7d6e;border-radius:12px;width:min(820px,94vw);max-width:94vw;max-height:86vh;padding:20px;overflow:auto}
#grow-ko-dialog::backdrop{background:rgba(15,23,42,.65)}#grow-ko-dialog h2{margin-right:0}#grow-ko-dialog h3{font-size:1.05rem}#grow-ko-dialog p,#grow-ko-dialog li{font-size:1rem;line-height:1.4}
.grow-ko-close{display:flex;justify-content:flex-end;margin:0 0 10px}.grow-ko-table-wrap{overflow-x:auto}.grow-ko-table-wrap table{width:100%;border-collapse:collapse}.grow-ko-table-wrap caption{text-align:left;font-weight:700;padding:8px 0}.grow-ko-table-wrap td,.grow-ko-table-wrap th{padding:8px;border:1px solid #bbb;text-align:left;overflow-wrap:anywhere}.grow-ko-table-wrap th{min-width:85px}
.grow-contact-model{margin:12px 0}.grow-contact-model svg{display:block;width:100%;height:auto}.grow-contact-model figcaption{font-size:.9rem;line-height:1.4}
.grow-exit-panel{margin-top:18px;border-top:2px solid #3f7d6e;padding-top:8px}.grow-exit-panel summary{cursor:pointer;font-weight:700;padding:12px 2px}.grow-exit-panel summary:focus-visible{outline:3px solid var(--ido-border);outline-offset:2px}
.grow-exit-panel fieldset{min-width:0;margin:14px 0;padding:12px;border:1px solid #778891;border-radius:8px}.grow-exit-panel legend{font-weight:700;max-width:100%;font-size:1rem;line-height:1.4}.grow-exit-choice{display:inline-flex;align-items:center;gap:6px;min-height:44px;margin-right:15px}.grow-exit-choice input{width:20px;height:20px}.grow-exit-panel textarea{display:block;width:100%;max-width:100%;min-height:76px;font:inherit;resize:vertical}.grow-exit-input-label{display:block;font-size:.95rem;margin-bottom:6px}.grow-exit-panel [hidden]{display:none!important}.grow-exit-explanations{margin-top:10px}.grow-exit-panel [aria-pressed="true"]{outline:2px solid currentColor;outline-offset:2px}
@media print{#grow-ko-dialog{display:none!important}.grow-ko-print{font-size:10pt;line-height:1.25}.grow-ko-print p,.grow-ko-print li,.grow-ko-print td,.grow-ko-print th{font-size:10pt!important;line-height:1.25!important}.grow-ko-print .grow-contact-model svg{max-height:36mm}.grow-ko-print h3{margin:5px 0!important}.grow-ko-print ul{margin:4px 0!important}.grow-ko-print .grow-ko-table-wrap td,.grow-ko-print .grow-ko-table-wrap th{padding:4px}.grow-exit-ticket{break-inside:avoid;border:1px solid #888;padding:5mm;margin-bottom:10mm;min-height:105mm}.grow-exit-ticket p{font-size:11pt;line-height:1.4}.grow-exit-writing-space{height:17mm}.grow-exit-ticket h3{font-size:14pt!important}}
</style>'''
 s=s.replace('</head>',css+'</head>')
 # Preserve total stage duration while naming the exit inside it.
 old='10 minutes. After sorting, test a changed case:'
 assert old in s;s=s.replace(old,'10 minutes total: about 8 for sorting, correction and the changed case; the final 2 for one Lesson A exit route. After sorting, test a changed case:',1)
 (root/REL).write_text(s)
 vector=fitz.open(stream=svg.encode(),filetype='svg')
 vector_pdf=fitz.open(stream=vector.convert_to_pdf(),filetype='pdf')
 png=vector_pdf[0].get_pixmap(matrix=fitz.Matrix(2,2)).tobytes('png')
 for who in ['Pupil','Teacher']:
  rel=PACK+f'GROW_W3A_Friction_{who}.docx';doc=Document(io.BytesIO(read(rel)))
  def heading(kicker,title):
   p=doc.add_paragraph(kicker,'Kicker');p.paragraph_format.page_break_before=True;doc.add_paragraph(title,'Title')
  def table_style(t):
   t.autofit=False
   for row in t.rows:
    for cell in row.cells:
     props=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
     for edge in ['top','left','bottom','right']:
      x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'5');x.set(qn('w:color'),'D9D9D9');b.append(x)
     props.append(b)
  if who=='Pupil':
   heading('KEEP NEARBY / KNOWLEDGE ORGANISER','Friction knowledge organiser')
   doc.add_paragraph('Learning objective: '+d['objective'])
   doc.add_paragraph('Success: '+' '.join(d['success']))
   t=doc.add_table(rows=1,cols=2);t.rows[0].cells[0].text='Word';t.rows[0].cells[1].text='Meaning'
   for w,v in d['vocabulary']:
    c=t.add_row().cells;c[0].text=w;c[1].text=v
   table_style(t)
   for fact in d['facts']:doc.add_paragraph(fact)
   doc.add_picture(io.BytesIO(png),width=Inches(5.8))
   doc.add_paragraph(d['model'])
   for example in d['examples']:doc.add_paragraph(example)
   doc.add_paragraph('Use evidence: '+d['comparison'])
   # Compact only the new reference page, preserving earlier reviewed pages.
   ko_start=len(doc.paragraphs)-len(d['facts'])-len(d['examples'])-4
   for p in doc.paragraphs[ko_start:]:
    p.paragraph_format.space_after=Pt(4)
    for r in p.runs:r.font.size=Pt(10)
   for k,r in d['routes'].items():
    heading('CHOOSE ONE ROUTE / TWO COPIES TO CUT','Friction exit '+r['label'])
    for repeat in range(2):
     doc.add_paragraph('Lesson A · '+r['label'],'Heading 1')
     doc.add_paragraph(d['exit_intro'])
     for i,q in enumerate(r['questions']):
      p=doc.add_paragraph(str(i+1)+'. '+q);p.runs[0].bold=True
      if k=='supported':doc.add_paragraph(' / '.join(r['options'][i]))
      else:
       p=doc.add_paragraph(' ');p.paragraph_format.space_after=Pt(20)
     p=doc.add_paragraph('Next lesson: '+d['next_lesson']);p.paragraph_format.space_after=Pt(18)
  else:
   doc.paragraphs[9].text='Each pupil: common pages 1–2, one response page (3–5), one arrival page (7–9), organiser page 10 and one exit route (11–13; two copies per sheet). Page 6 has reusable cards. Use screen or paper once; do not require both.'
   doc.tables[0].rows[-1].cells[2].text='About 8 min: sort, justify and challenge a changed case. Final 2 min: one Lesson A exit route. Label/save and stop at 40 minutes.'
   heading('TEACHER ANSWERS / LESSON A EXIT','Friction exit guidance')
   doc.add_paragraph(d['staff'])
   for k,r in d['routes'].items():
    doc.add_paragraph(r['label'],'Heading 1')
    for i,a in enumerate(r['answers']):doc.add_paragraph(str(i+1)+'. '+a)
   doc.add_paragraph('The shared knowledge organiser is on pupil page 10. It names contacts, qualifies roughness and energy claims, and separates the model from actual observations.')
   doc.add_paragraph('Next lesson: '+d['next_lesson'])
  doc.save(root/rel)
if __name__=='__main__':
 ap=argparse.ArgumentParser();ap.add_argument('--root',type=Path,default=HERE.parents[1]);a=ap.parse_args();main(a.root.resolve())
